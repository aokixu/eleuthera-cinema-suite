"""Dossier Excel: hojas filtrables y plantilla de producción lista para imprimir."""
from ui_i18n import ui_text, ui_join
from pathlib import Path
from datetime import datetime
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from project_dossier import _scene_labels, _readable, LABELS

LABELS = {**LABELS, 'text':'Texto', 'type':ui_text('Tipo'), 'name':ui_text('Nombre'), 'section':'Sección',
          'field':'Campo', 'value':'Contenido', 'character':ui_text('Personaje'), 'beat':ui_text('Hito'),
          'model':'Modelo', 'planning':'Planificación', 'purpose':'Función dramática',
          'transformation':'Transformación', 'objective':ui_text('Objetivo'), 'conflict':ui_text('Conflicto'),
          'value_start':'Valor inicial', 'value_end':'Valor final', 'part':'Parte',
          'shooting_days':'Días de rodaje', 'base_currency':ui_text('Moneda base'),
          'default_fringe':'Carga predeterminada (%)'}
NAVY='172B36'; TEAL='17616A'; GOLD='D8B565'; INK='233A46'
NUMERIC={'quantity','days','unit_price','exchange','fringe','total','pages','order','day'}
LONG={'text','notes','source','synopsis','description','planning','value','need','initial_state','final_state'}
TECHNICAL={'id','x','y','scene_id'}


def _plain(value):
    if value is None: return ''
    if isinstance(value,bool): return 'Sí' if value else 'No'
    if isinstance(value,dict): return '\n'.join(f'{LABELS.get(k,k)}: {_plain(v)}' for k,v in value.items())
    if isinstance(value,list): return '\n'.join(_plain(v) for v in value)
    return str(value)


def _put(cell, value):
    cell.value=value
    # Project text must remain literal, including strings beginning with '='.
    if isinstance(value,str): cell.data_type='s'


def _fields(value, prefix=''):
    if isinstance(value,dict):
        for key,item in value.items():
            if key in TECHNICAL or key.endswith('_ids'): continue
            label=LABELS.get(key,key.replace('_',' ').capitalize())
            yield from _fields(item, f'{prefix} / {label}' if prefix else label)
    elif isinstance(value,list):
        for i,item in enumerate(value,1): yield from _fields(item,f'{prefix} / {i}')
    elif value is not None and str(value).strip():
        yield {'field':prefix,'value':_plain(value)}


def _chunks(value):
    text = _plain(value)
    if not text: return ['']
    result=[]
    while text:
        end=min(500,len(text))
        # Excel's maximum row height is finite; split newline-heavy text as well.
        newlines=[i for i,c in enumerate(text[:end]) if c=='\n']
        if len(newlines)>18: end=newlines[17]+1
        result.append(text[:end]);text=text[end:]
    return result


def _sheet(book, name, title, rows, columns=None):
    rows=list(rows)
    keys=list(columns or [])
    for row in rows:
        for key in row:
            if key not in keys and key not in TECHNICAL and not key.endswith('_ids'):
                keys.append(key)
    if not keys: keys=['value']
    # Keep long text visible rather than silently clipping it at Excel's row/cell limits.
    expanded=[]
    for row in rows:
        chunks={k: _chunks(row.get(k,'')) for k in keys}
        parts=max(len(v) for v in chunks.values())
        for part in range(parts):
            item={k: (v[part] if part<len(v) else (v[0] if k in ('name','character','title','code','order') else '')) for k,v in chunks.items()}
            if parts>1: item['part']=f'{part+1}/{parts}'
            expanded.append(item)
    if any('part' in row for row in expanded) and 'part' not in keys: keys.append('part')
    ws=book.create_sheet(name); n=len(keys); end=get_column_letter(max(2,n))
    ws.sheet_properties.tabColor=TEAL;ws.sheet_view.showGridLines=False
    ws.merge_cells(f'A1:{end}1');_put(ws['A1'],title);ws['A1'].font=Font(name='Aptos Display',size=20,bold=True,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor=NAVY);ws['A1'].alignment=Alignment(vertical='center');ws.row_dimensions[1].height=46
    ws.merge_cells(f'A2:{end}2');_put(ws['A2'],f'ELEUTHERA  /  {name.upper()}  /  {len(rows)} registros');ws['A2'].font=Font(name='Aptos',size=10,bold=True,color=TEAL);ws.row_dimensions[2].height=26
    ws.merge_cells(f'A3:{end}3');_put(ws['A3'],'← Volver al resumen');ws['A3'].hyperlink="#'Resumen'!A1";ws['A3'].font=Font(name='Aptos',size=10,color=TEAL,underline='single')
    widths={}
    for col,key in enumerate(keys,1):
        letter=get_column_letter(col); width=65 if key in LONG else (40 if key in ('scene_label','heading','concept','field','location','title','elements') else 22)
        if key in NUMERIC: width=13
        if key in ('code','level','unit','block','category','currency','part'): width=16
        if name==ui_text('Presupuesto') and key in ('concept','notes'): width=36
        widths[key]=width;ws.column_dimensions[letter].width=width
        cell=ws.cell(5,col);_put(cell,LABELS.get(key,key.replace('_',' ').capitalize()));cell.font=Font(name='Aptos',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor=TEAL);cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[5].height=30
    for r,row in enumerate(expanded,6):
        height=24
        for col,key in enumerate(keys,1):
            value=row.get(key,'');cell=ws.cell(r,col)
            if key in NUMERIC and value:
                try:
                    number=float(value)
                    if math.isfinite(number): value=number
                except (ValueError,TypeError): pass
            _put(cell,value);cell.font=Font(name='Aptos',size=11,color=INK)
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            cell.fill=PatternFill('solid',fgColor='F0F5F6' if r%2==0 else 'FFFFFF')
            cell.border=Border(bottom=Side(style='hair',color='DCE5E8'))
            if key in NUMERIC: cell.number_format='#,##0.00;[Red](#,##0.00);–'
            lines=sum(max(1,math.ceil(len(line)/max(8,int(widths[key]*0.85)))) for line in str(value).split('\n'))
            height=max(height,lines*16+10)
        ws.row_dimensions[r].height=min(409,height)
    if expanded:
        table=Table(displayName=f'Seccion{len(book.worksheets)}',ref=f'A5:{get_column_letter(n)}{ws.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    else:
        _put(ws.cell(6,1),'Sin datos registrados.')
    ws.freeze_panes='C6' if n>3 else 'A6'
    ws.print_title_rows='1:5';ws.print_options.horizontalCentered=True
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape' if n>3 else 'portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A3 if n>8 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_area=ws.dimensions;ws.oddFooter.left.text='Eleuthera | '+name;ws.oddFooter.right.text='Página &P de &N'
    return ws


def write_master_excel(filename,data,blocks):
    scenes=_scene_labels({'blocks':blocks});data=_readable(data,scenes)
    info=data.get('project_info') or {};title=info.get('title') or data.get('title') or 'Proyecto Eleuthera'
    book=Workbook();book.remove(book.active);book.properties.creator='Eleuthera Cinema Suite';book.properties.title=title
    cover=book.create_sheet('Resumen');cover.sheet_view.showGridLines=False;cover.sheet_properties.tabColor=GOLD
    for col,width in {'A':4,'B':34,'C':28,'D':28,'E':28,'F':4}.items():cover.column_dimensions[col].width=width
    cover.merge_cells('B2:E2');_put(cover['B2'],'ELEUTHERA / DOSSIER DE PRODUCCIÓN');cover['B2'].font=Font(name='Aptos',size=12,bold=True,color=TEAL)
    cover.merge_cells('B4:E6');_put(cover['B4'],title);cover['B4'].font=Font(name='Aptos Display',size=26,bold=True,color='FFFFFF');cover['B4'].fill=PatternFill('solid',fgColor=NAVY);cover['B4'].alignment=Alignment(wrap_text=True,vertical='center');cover.row_dimensions[4].height=42
    cover.merge_cells('B8:E8');_put(cover['B8'],f'Exportado {datetime.now():%d/%m/%Y %H:%M} · Versión {info.get("script_version", "—")}');cover['B8'].font=Font(name='Aptos',size=11,color=INK)
    for row,(label,value) in enumerate([(ui_text('Productora'),info.get('company','')),(ui_text('Dirección'),info.get('director','')),(ui_text('Guion'),info.get('writer','')),(ui_text('Escenas'),len(scenes)),('Partidas de presupuesto',len(data.get('budget') or [])),(ui_text('Moneda base'),(data.get('budget_settings') or {}).get('base_currency',info.get('currency','')))],10):
        _put(cover.cell(row,2),label);cover.cell(row,2).font=Font(name='Aptos',bold=True,color=TEAL)
        cover.merge_cells(start_row=row,start_column=3,end_row=row,end_column=5);_put(cover.cell(row,3),value);cover.row_dimensions[row].height=26
    _sheet(book,ui_text('Desarrollo'),title,_fields(data.get('treatment') or {}),['field','value'])
    characters=[];milestones=[]
    for name,arc in (data.get('character_arcs') or {}).items():
        if not isinstance(arc,dict):arc={'notes':arc}
        characters.append({'name':name,**{k:v for k,v in arc.items() if k!='milestones'}})
        milestones.extend({'character':name,**r} for r in arc.get('milestones',[]) if isinstance(r,dict))
    _sheet(book,ui_text('Personajes'),title,characters,['name','initial_state','need','final_state'])
    _sheet(book,'Hitos personajes',title,milestones,['character','moment','title','scene_label','notes'])
    structure=data.get('dramatic_structure') or {};model=structure.get('structure_model','');points=structure.get('structure_points') or {};planned=structure.get('structure_planned_scenes') or {};planned=planned.get(model,planned);dev=(structure.get('structure_development') or {}).get(model,{})
    rows=[]
    for name in dict.fromkeys([*points,*planned,*dev]):
        rows.append({'beat':name,'model':model,'scene_label':scenes.get(str(points.get(name,'')),'Sin escena vinculada'),'planning':planned.get(name,''),**(dev.get(name,{}) if isinstance(dev.get(name),dict) else {})})
    _sheet(book,'Estructura',title,rows,['beat','model','scene_label','planning'])
    types={'scene':ui_text('Escena'),'action':ui_text('Acción'),'character':ui_text('Personaje'),'dialogue':ui_text('Diálogo'),'parenthetical':'Paréntesis','transition':ui_text('Transición')}
    _sheet(book,ui_text('Guion'),title,({'order':i,'type':types.get(b.get('type'),b.get('type','')),'text':b.get('text','')} for i,b in enumerate(blocks,1)),['order','type','text'])
    _sheet(book,ui_text('Desglose'),title,data.get('breakdown') or [],['scene_label','category','element','approved','source'])
    _sheet(book,'Rodaje',title,data.get('schedule') or [],['order','day','date','scene','location','period','pages','synopsis','elements','status'])
    _sheet(book,ui_text('Presupuesto'),title,data.get('budget') or [],['code','level','block','category','concept','quantity','unit','days','unit_price','currency','exchange','fringe','total','notes'])
    _sheet(book,ui_text('Mapa de tramas'),title,data.get('story_map') or [],['title','lane','scene_label'])
    _sheet(book,ui_text('Datos del proyecto'),title,_fields(info),['field','value'])
    settings=[]
    for key in ('budget_settings','budget_globals','budget_fringes','budget_groups','budget_credits','budget_charges'):
        settings.extend(_fields(data.get(key) or {},key.replace('_',' ')))
    _sheet(book,'Ajustes presupuesto',title,settings,['field','value'])
    cover.merge_cells('B18:E18');_put(cover['B18'],'CONTENIDO · Haz clic para abrir una sección');cover['B18'].font=Font(name='Aptos',bold=True,color=TEAL,size=12)
    for row,ws in enumerate(book.worksheets[1:],20):
        cover.merge_cells(start_row=row,start_column=2,end_row=row,end_column=5);cell=cover.cell(row,2);_put(cell,f'{row-19:02d}   {ws.title}');cell.hyperlink=f"#'{ws.title}'!A1";cell.font=Font(name='Aptos',size=12,color=TEAL,underline='single');cell.fill=PatternFill('solid',fgColor='F0F5F6');cover.row_dimensions[row].height=28
    cover.print_area=f'B2:E{cover.max_row}';cover.sheet_properties.pageSetUpPr.fitToPage=True;cover.page_setup.fitToWidth=1;cover.page_setup.fitToHeight=1
    book.save(filename)


def export_master_excel(parent,data,blocks):
    from PySide6.QtWidgets import QFileDialog,QMessageBox
    filename,_=QFileDialog.getSaveFileName(parent,ui_text('Exportar dossier Excel'),'','Excel (*.xlsx)')
    if not filename:return
    if not filename.lower().endswith('.xlsx'):filename+='.xlsx'
    try:write_master_excel(filename,data,blocks)
    except (OSError,ValueError) as error:
        QMessageBox.critical(parent,ui_text('No se pudo exportar'),str(error));return
    QMessageBox.information(parent,ui_text('Dossier Excel'),ui_text('Proyecto exportado con portada y hojas por sección.'))
