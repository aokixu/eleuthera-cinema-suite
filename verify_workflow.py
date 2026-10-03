"""Integration checks: stable identities, recovery, conservative sync and reports."""
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
from PySide6.QtCore import QSettings, Qt, QSize
from PySide6.QtGui import QTextCursor
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox, QFileDialog, QPlainTextEdit
from PySide6.QtTest import QTest
from PySide6.QtPdf import QPdfDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from suite import SuiteWindow
from project_store import atomic_json, fingerprint, snapshots
from production_reports import write_pdf, write_excel, breakdown_tables, budget_tables, schedule_tables
from breakdown_export import export_breakdown_workbook

qt = QApplication([])
from PySide6.QtGui import QFontDatabase, QPainter, QImage
for font in ('segoeui.ttf', 'segoeuib.ttf', 'arial.ttf', 'arialbd.ttf', 'cour.ttf', 'courbd.ttf'):
    QFontDatabase.addApplicationFont('C:/Windows/Fonts/' + font)
qt.setOrganizationName('EleutheraWorkflowTests'); qt.setApplicationName('WorkflowTests'); qt.setStyle('Fusion')
QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Ok)
QMessageBox.warning = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Ok)
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
def fail(*args, **kwargs): raise AssertionError(args[-1])
QMessageBox.critical = staticmethod(fail)
artifacts = Path('workflow-checks'); artifacts.mkdir(exist_ok=True)
with TemporaryDirectory() as temp:
    temp = Path(temp)
    window = SuiteWindow('professional'); window.poll.stop(); window.autosave.stop(); window.recovery_root = temp / 'recovery'
    window.resize(1480, 950); window.show(); qt.processEvents()
    blocks = [
        {'type': 'scene', 'text': 'INT. HOSPITAL - NOCHE'},
        {'type': 'action', 'text': 'Pedro mira el reloj. Suena un teléfono.'},
        {'type': 'character', 'text': 'PEDRO'},
        {'type': 'dialogue', 'text': 'Busco a Elena.'},
        {'type': 'scene', 'text': 'INT. HOSPITAL - NOCHE'},
        {'type': 'action', 'text': 'Una fotografía sobre la mesa.'},
        {'type': 'character', 'text': 'ELENA'},
    ]
    window.editor.load_blocks(blocks); window.analyze_script()
    scenes = window.current_scenes(); ids = [s['id'] for s in scenes]
    assert len(set(ids)) == 2
    assert {r['scene_id'] for r in window.breakdown_data()} == set(ids)
    # Renaming a scene must retain its ID, corrections and approval.
    window.breakdown.item(0, 0).setCheckState(Qt.CheckState.Checked)
    window.breakdown.item(0, 3).setText('Localización corregida manualmente')
    cursor = QTextCursor(window.editor.document().firstBlock()); cursor.select(QTextCursor.SelectionType.BlockUnderCursor); cursor.insertText('INT. HOSPITAL CENTRAL - NOCHE')
    assert window.current_scenes()[0]['id'] == ids[0]
    assert any(c['id'] == ids[0] and c['module'] == 'breakdown' for c in window.pending_changes())
    window.choose_rows = lambda *a, **k: []
    before = window.breakdown_data(); window.analyze_script(); assert window.breakdown_data() == before
    window.sync_schedule()
    assert len(window.schedule_page.to_data()) == 2
    assert all(r['day'] == '' for r in window.schedule_page.to_data())
    assert sum(i[0] == 'Escena sin jornada' for i in window.export_issues('schedule')) == 2
    window.schedule_page.table.item(0, 1).setText('4')
    window.schedule_page.table.item(0, 8).setText('Sinopsis manual')
    window.schedule_page.table.selectRow(1); window.schedule_page.move_selected(-1)
    old_order = [r['scene_id'] for r in window.schedule_page.to_data()]
    window.breakdown.item(0, 3).setText('Nueva necesidad manual')
    assert any(c['module'] == 'schedule' for c in window.pending_changes())
    old_plan = window.schedule_page.to_data(); window.sync_schedule(); assert window.schedule_page.to_data() == old_plan
    window.choose_rows = lambda title, records, describe, default=False: records
    window.sync_schedule()
    assert [r['scene_id'] for r in window.schedule_page.to_data()] == old_order
    assert next(r for r in window.schedule_page.to_data() if r['scene_id'] == ids[0])['day'] == '4'
    assert not hasattr(window, 'search_page')
    # Metadata cannot overwrite the authoritative Budget currency catalog.
    window.add_budget_row(concept='Camara', data={'concept': 'Camara', 'quantity': '2', 'days': '3', 'unit_price': '100', 'currency': 'USD', 'exchange': '1'})
    window.info_fields['title'].setText('Pel\u00edcula de prueba')
    import budget_currencies as money
    window.budget_currencies = {'version': 1, 'base': 'BOB', 'currencies': [
        {'code': 'BOB', 'symbol': 'Bs', 'rate': 1.0}, {'code': 'USD', 'symbol': '$', 'rate': 6.96}]}
    money.mirror(window); window.update_budget_total()
    window.info_fields['currency'].setText('EUR'); window.rates.setText('6.96')
    window.apply_metadata(); assert window.base_currency.currentText() == 'BOB'
    assert window.budget.item(0, 12).text() == '4176.00'
    window.add_budget_row(concept='Sin tarifa')
    assert any(i[0] == 'Partida sin tarifa' for i in window.export_issues('budget'))
    window.budget.item(1, 8).setText('0')
    assert not any(i[0] == 'Partida sin tarifa' for i in window.export_issues('budget'))
    # Editing really occupies the visible cell, commits, cancels, and protects totals.
    window.modules.setCurrentWidget(window.budget_page); window.budget.setCurrentCell(0, 4); qt.processEvents()
    window.budget.editItem(window.budget.item(0, 4)); qt.processEvents()
    editor = next(e for e in window.budget.findChildren(QPlainTextEdit) if e.isVisible())
    rect = window.budget.visualItemRect(window.budget.item(0, 4))
    assert editor.width() >= rect.width() - 4 and editor.height() >= rect.height() - 4
    editor.setPlainText('Cámara principal'); QTest.keyClick(editor, Qt.Key.Key_Return); qt.processEvents()
    assert window.budget.item(0, 4).text() == 'Cámara principal'
    assert not window.budget.item(0, 12).flags() & Qt.ItemFlag.ItemIsEditable
    window.budget.editItem(window.budget.item(0, 4)); qt.processEvents()
    editor = next(e for e in window.budget.findChildren(QPlainTextEdit) if e.isVisible()); editor.setPlainText('No guardar'); QTest.keyClick(editor, Qt.Key.Key_Escape); qt.processEvents()
    assert window.budget.item(0, 4).text() == 'Cámara principal'
    # Unnamed sessions recover all modules; keep only five versions.
    window.path = None
    for version in range(7):
        window.notes.setPlainText('Versión ' + str(version)); window.auto_save()
    candidates = window.recovery_candidates(); assert len(candidates) == 5
    recovered = candidates[0][1]['data']; assert recovered['notes'] == 'Versión 6'
    assert recovered['budget'][0]['total'] == '4176.00' and recovered['project_info']['title'] == 'Película de prueba'
    assert recovered['schedule'][0]['scene_id'] in ids
    window.path = temp / 'movie.eguion'; assert window.save()
    for version in range(7): window.notes.setPlainText('Guardado ' + str(version)); assert window.save()
    assert len(list((temp / '.eleuthera-backups').glob('*/*.eguion'))) == 5
    saved = json.loads(window.path.read_text(encoding='utf-8'))
    window.load_project(saved, window.path); assert [s['id'] for s in window.current_scenes()] == ids
    assert window.project_info['title'] == 'Película de prueba'
    assert window.new_document(); assert not window.breakdown_data() and not window.budget_data() and not window.schedule_page.to_data()
    window.load_project(saved, None)
    # Legacy ambiguous headings are not merged or guessed.
    legacy = {'blocks': blocks, 'breakdown': [{'approved': False, 'scene': blocks[0]['text'], 'category': 'Utilería', 'element': 'Reloj', 'source': 'Texto'}]}
    window.load_project(legacy, None); assert window.breakdown_data()[0]['scene_id'] == ''
    assert any(c['module'] == 'legacy' for c in window.pending_changes())
    window.load_project(saved, None)
    rows = window.breakdown_data()
    for row in rows: row['scene'] = 'INT. HOSPITAL - NOCHE'
    export_breakdown_workbook(artifacts / 'desglose.xlsx', rows, 'Película de prueba', window.project_info)
    workbook = load_workbook(artifacts / 'desglose.xlsx'); assert workbook.worksheets[0].max_row == 5
    write_pdf(artifacts / 'desglose.pdf', 'Película de prueba', window.project_info, breakdown_tables(rows))
    write_pdf(artifacts / 'presupuesto.pdf', 'Película de prueba', window.project_info, budget_tables(window.budget_data(), window.total_label.text()))
    write_excel(artifacts / 'presupuesto.xlsx', 'Película de prueba', window.project_info, budget_tables(window.budget_data(), window.total_label.text()))
    write_pdf(artifacts / 'plan.pdf', 'Película de prueba', window.project_info, schedule_tables(window.schedule_page.to_data()))
    for name in ('desglose', 'presupuesto', 'plan'):
        pdf = PdfReader(artifacts / (name + '.pdf')); assert len(pdf.pages) >= 1
        text = '\n'.join(p.extract_text() for p in pdf.pages); assert 'prueba' in text.lower()
        doc = QPdfDocument(); doc.load(str(artifacts / (name + '.pdf'))); render = doc.render(0, QSize(1500, 1060))
        canvas = QImage(render.size(), QImage.Format.Format_RGB32); canvas.fill(Qt.GlobalColor.white)
        painter = QPainter(canvas); painter.drawImage(0, 0, render); painter.end(); canvas.save(str(artifacts / (name + '-preview.png')))
    window.modules.setCurrentWidget(window.budget_page); qt.processEvents(); window.grab().save(str(artifacts / 'professional.png'))
    window.budget.editItem(window.budget.item(0, 4)); qt.processEvents(); window.grab().save(str(artifacts / 'cell-editor.png'))
    window.commit_active_cell(); window._last_saved = fingerprint(window.project_data()); window.dirty = False; window.close()
    standard = SuiteWindow('standard'); standard.poll.stop(); standard.autosave.stop(); standard.recovery_root = temp / 'standard-recovery'; standard.show(); qt.processEvents()
    assert not hasattr(standard, 'excel_button')
    assert not standard.modules.isTabVisible(standard.modules.indexOf(standard.schedule_page))
    standard.modules.setCurrentWidget(standard.budget_page); qt.processEvents(); standard.grab().save(str(artifacts / 'standard.png'))
    standard._last_saved = fingerprint(standard.project_data()); standard.dirty = False; standard.close()
print('OK: stable scenes, conservative sync, metadata, cell editing, recovery, backups, legacy migration, PDF/Excel, both editions.')
