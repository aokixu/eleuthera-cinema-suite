import copy,json
from pathlib import Path
from time import perf_counter
from unittest.mock import patch
from decimal import Decimal
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication,QDialog,QLineEdit,QDialogButtonBox,QMessageBox,QFileDialog
from openpyxl import load_workbook
from suite import SuiteWindow
from project_store import fingerprint
from budget_values import BudgetError
from budget_groups_ui import GroupsDialog,GroupSelectionDialog,GroupSummaryDialog,export_groups
from budget_group_analysis import analyze

app=QApplication([]);out=Path('test-results/stage4');out.mkdir(exist_ok=True)
metrics,passed,warnings={},[],[]
QMessageBox.warning=lambda *a,**k:warnings.append(a)
QMessageBox.information=lambda *a,**k:None
QMessageBox.question=lambda *a,**k:QMessageBox.StandardButton.Yes
QMessageBox.critical=lambda *a,**k:(_ for _ in ()).throw(AssertionError(a))
def window():
    w=SuiteWindow('professional');w.poll.stop();w.autosave.stop();return w
def drain(predicate):
    deadline=perf_counter()+60
    while predicate() and perf_counter()<deadline:app.processEvents()
    assert not predicate()
def close(w):w._last_saved=fingerprint(w.project_data());w.dirty=False;w.close()
def rejects(w,call):
    before=fingerprint(w.project_data())
    try:call()
    except BudgetError:pass
    else:raise AssertionError('Invalid operation accepted')
    assert fingerprint(w.project_data())==before

w=window();c=w.budget_controller;g=c.groups
dialog=GroupsDialog(g,w);old_exec=QDialog.exec;fields=('UNIDAD_A','Primera unidad')
def fill(form):
    for field,text in zip(form.findChildren(QLineEdit),fields):field.setText(text)
    form.findChild(QDialogButtonBox).accepted.emit();return form.result()
QDialog.exec=fill;dialog.edit(True);a=g.catalog.items()[0].id
before=fingerprint(w.project_data());dialog.edit(True);assert warnings and fingerprint(w.project_data())==before
fields=('UNIDAD_A_RENOMBRADA','Nueva descripción');dialog.edit(False);assert g.catalog.get(a).description=='Nueva descripción'
QDialog.exec=old_exec
b=g.mutate('add',name='UNIDAD_B');unused=g.mutate('add',name='TEMPORAL')
dialog.table.selectRow(2);dialog.remove();assert len(g.catalog.items())==2
c.mutate('add','TARIFA','100','');fringe=c.mutate_fringe('add',name='SEGURO',percentage='10')
w.add_budget_row(data={'concept':'Inicial','quantity':'1','days':'1','unit_price':'TARIFA','fringe_ids':[fringe]})
selection=GroupSelectionDialog(g,0,w)
for i in range(selection.listing.count()):selection.listing.item(i).setCheckState(Qt.CheckState.Checked)
selection.save();assert g.analysis.totals[a]==110 and g.analysis.totals[b]==110
dialog.table.selectRow(0);dialog.toggle();assert not g.catalog.get(a).active and w.budget_data()[0]['group_ids']==[a,b]
dialog.toggle();dialog.remove();assert 'fila 1' in str(warnings[-1])
dialog.show();app.processEvents();dialog.grab().save(str(out/'groups.png'));dialog.close()
passed.append('UI crear/editar/renombrar/activar/eliminar, asignación múltiple y eliminación en uso')
rejects(w,lambda:g.assign([0],['missing']))
rejects(w,lambda:g.assign([0,999],[a]))
rejects(w,lambda:g.mutate('edit',a,name='UNIDAD_B'))
g.assign([0],[a,a]);assert w.budget_data()[0]['group_ids']==[a]
before=fingerprint(w.project_data());w.budget.item(0,0).setText('Cuenta')
assert w.budget.item(0,0).text()=='Detalle' and fingerprint(w.project_data())==before
passed.append('Duplicados y asignaciones inválidas: transacciones sin cambios parciales')

w._loading=True;w.budget.setRowCount(0)
w.add_budget_row(data={'level':'Cuenta','concept':'Equipo','block':'ATL'})
for i in range(1025):
    w.add_budget_row(data={'concept':'Partida '+str(i),'block':'ATL','quantity':'1','days':'1','unit_price':'TARIFA',
                          'fringe_ids':[fringe],'group_ids':[a] if i<25 else [b] if i<55 else []})
w._loading=False;c.request();drain(lambda:c.busy)
assert len(g.analysis.members[a])==25 and len(g.analysis.members[b])==30
assert g.analysis.totals[a]==2750 and g.analysis.totals[b]==3300
start=perf_counter();w.budget.item(1,5).setText('2');drain(lambda:c.busy)
metrics['edit_assigned_row_ms']=(perf_counter()-start)*1000
assert c.last_run['evaluated_rows']==1 and g.analysis.visited_rows==1 and g.analysis.changed=={a}
assert g.analysis.totals[a]==2860 and g.analysis.totals[b]==3300
metrics['row_evaluated']=c.last_run['evaluated_rows'];metrics['groups_updated']=len(g.analysis.changed)
with patch('budget_engine.row_total',side_effect=AssertionError('Unnecessary budget recalculation')):
    g.assign([1],[a,b]);assert g.analysis.totals[b]==3520
    g.assign([1],[a]);assert g.analysis.totals[b]==3300
passed.append('1025 partidas: 25 A y 30 B; editar una evalúa una fila y actualiza solo A; reasignar no evalúa presupuesto')

before=fingerprint(w.project_data());start=perf_counter();g.set_filter([a,b])
metrics['filter_1026_rows_ms']=(perf_counter()-start)*1000
assert sum(not w.budget.isRowHidden(i) for i in range(w.budget.rowCount()))==56
assert fingerprint(w.project_data())==before
g.set_filter([a]);assert w.budget.isRowHidden(26) and not w.budget.isRowHidden(0)
w.show_all_groups_button.click();assert not any(w.budget.isRowHidden(i) for i in range(w.budget.rowCount()))
g.set_filter([a]);g.assign([26],[a,b]);assert not w.budget.isRowHidden(26)
g.assign([26],[b]);assert w.budget.isRowHidden(26)
passed.append('Filtro unión, contexto de Cuenta, Mostrar todo y cambios de pertenencia sin alterar datos')

start=perf_counter();summary=analyze(c.rows,c.result.totals,g.catalog)
metrics['aggregate_1026_rows_ms']=(perf_counter()-start)*1000
assert summary.totals==g.analysis.totals
project=out/'groups-project.eguion';w.path=project
start=perf_counter();assert w.save();metrics['save_1026_rows_ms']=(perf_counter()-start)*1000
saved=w.project_data();assert len(saved['budget'])==1026
close(w);w=window();c=w.budget_controller;g=c.groups
start=perf_counter();assert w.open_path(project);drain(lambda:getattr(w,'_import_job',None) is not None)
metrics['reopen_1026_rows_ms']=(perf_counter()-start)*1000
assert w.project_data()['budget_groups']==saved['budget_groups'] and g.analysis.totals[a]==2860
assert not g.selected_filter and not any(w.budget.isRowHidden(i) for i in range(w.budget.rowCount()))
for ids in (['missing'],None):
    corrupt=copy.deepcopy(saved);corrupt['budget'][1]['group_ids']=ids
    rejects(w,lambda:w.load_project(corrupt))
passed.append('Guardar filtrado conserva todas las filas; cerrar/reabrir reconstruye Groups; referencias corruptas rechazadas')

report=out/'groups-summary.xlsx';QFileDialog.getSaveFileName=lambda *a,**k:(str(report),'')
view=GroupSummaryDialog(g,w)
c.mutate('edit','TARIFA','TARIFA','200','');assert c.busy
assert view.table.item(0,2).text()=='Pendiente'
pending=out/'pending.xlsx';QFileDialog.getSaveFileName=lambda *a,**k:(str(pending),'')
export_groups(g,w);assert not pending.exists()
drain(lambda:c.busy);assert g.analysis.totals[a]==5720 and g.analysis.totals[b]==6600
c.mutate_fringe('edit',fringe,name='SEGURO',percentage='20');drain(lambda:c.busy)
assert g.analysis.totals[a]==6240
g.set_filter([a]);QFileDialog.getSaveFileName=lambda *a,**k:(str(report),'')
start=perf_counter();export_groups(g,w);metrics['export_summary_ms']=(perf_counter()-start)*1000
book=load_workbook(report);data=list(book.active.iter_rows(values_only=True))
assert any(row[0]=='UNIDAD_A_RENOMBRADA' and row[1]==25 and row[2]=='6240.00' for row in data)
assert any(row[0]=='UNIDAD_B' and row[1]==30 and row[2]=='7200.00' for row in data)
view.show();app.processEvents();view.grab().save(str(out/'groups-summary.png'));view.close()
passed.append('Globals y Fringes actualizan Groups; resumen pendiente protegido; Excel usa totales completos aunque haya filtro')

# Real currency migration: totals already converted, never converted again by Groups.
small=copy.deepcopy(saved);small['budget']=small['budget'][1:2]
small['budget'][0].update(quantity='1',unit_price='100',exchange='6.96',currency='USD')
small.pop('budget_currencies',None);small.pop('budget_currency_legacy',None)
small['project_info'].update(currency='BOB',rates={'USD':6.96});small['budget_settings']['base_currency']='BOB'
w.load_project(small);assert g.analysis.totals[a]==Decimal('765.6')
w.budget.item(0,8).setText('1 / 0');assert c.result.errors
rejects(w,lambda:g.summary_rows())
w.budget.item(0,8).setText('100');assert not c.result.errors and g.analysis.totals[a]==Decimal('765.6')
small.pop('budget_groups');small['budget'][0].pop('group_ids');w.load_project(small)
assert g.catalog.items()==() and not g.selected_filter
passed.append('Moneda + Fringe + Groups sin doble conversión; error de partida recuperable; proyecto antiguo sin Groups')
close(w)
(out/'integration-results.json').write_text(json.dumps(dict(passed=passed,metrics=metrics),ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS:',len(passed),'escenarios de Groups');print(json.dumps(metrics,indent=2))
