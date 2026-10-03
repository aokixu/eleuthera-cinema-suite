"""Readable production reports using the application's existing Qt/openpyxl stack."""
from ui_i18n import ui_text, ui_join
from collections import OrderedDict, defaultdict
from html import escape
from pathlib import Path
from math import ceil
from PySide6.QtCore import QSizeF
from PySide6.QtGui import QFont, QPageLayout, QPageSize, QTextDocument
from PySide6.QtPrintSupport import QPrinter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LABELS = {'company': ui_text('Productora'), 'director': ui_text('Dirección'), 'producer': ui_text('Producción'), 'writer': ui_text('Guion'), 'script_version': ui_text('Versión del guion'), 'script_date': 'Revisión', 'currency': ui_text('Moneda base'), 'production_start': ui_text('Inicio de producción'), 'production_end': ui_text('Fin de producción')}


def metadata_lines(info):
    return [label + ': ' + str(info[key]) for key, label in LABELS.items() if info.get(key)]


def text_cell(cell, value):
    cell.value = value
    if isinstance(value, str): cell.data_type = 's'


def write_excel(path, title, info, tables):
    book = Workbook(); book.remove(book.active)
    for name, headers, rows in tables:
        sheet = book.create_sheet(name[:31]); n = len(headers)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        text_cell(sheet.cell(1, 1), title + ' · ' + name)
        sheet.cell(1, 1).font = Font(size=16, bold=True, color='FFFFFF'); sheet.cell(1, 1).fill = PatternFill('solid', fgColor='7B2638'); sheet.row_dimensions[1].height = 36
        metadata = metadata_lines(info)
        for rownum, value in enumerate(metadata, 2):
            sheet.merge_cells(start_row=rownum, start_column=1, end_row=rownum, end_column=n); text_cell(sheet.cell(rownum, 1), value)
        top = len(metadata) + 3
        for col, value in enumerate(headers, 1):
            cell = sheet.cell(top, col, value); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='57534E')
        for rownum, values in enumerate(rows, top + 1):
            for col, value in enumerate(values, 1):
                cell = sheet.cell(rownum, col); text_cell(cell, value)
                cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.font = Font(size=11)
                if rownum % 2 == 0: cell.fill = PatternFill('solid', fgColor='F5F1ED')
            sheet.row_dimensions[rownum].height = min(400, max(32, 15 * max((sum(max(1, ceil(len(p) / 36)) for p in str(v).split('\n')) for v in values), default=1)))
        for col, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(col)].width = 42 if any(w in header.lower() for w in (ui_text('escena'), 'concepto', 'elemento', 'notas', 'sinopsis', 'local')) else 18
        sheet.freeze_panes = f'A{top + 1}'; sheet.auto_filter.ref = f'A{top}:{get_column_letter(n)}{max(top, sheet.max_row)}'
        sheet.sheet_view.showGridLines = False; sheet.print_title_rows = f'1:{top}'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0; sheet.print_area = sheet.dimensions
        sheet.oddFooter.center.text = 'Eleuthera · Página &P de &N'
    book.save(path)


def report_html(title, info, tables):
    def html(value): return escape(str(value)).replace('\n', '<br>')
    content = ['<html><head><style>body {font-family: Arial; font-size:9pt; color:#292524;} h1 {color:#7b2638;font-size:17pt;} h2 {font-size:12pt;} table {border-collapse:collapse;width:100%;} th {background:#57534e;color:white;} td,th {padding:5px;border:1px solid #d8d2cc;} </style></head><body>', '<h1>' + html(title) + '</h1>', '<p>' + '<br>'.join(html(line) for line in metadata_lines(info)) + '</p>']
    for name, headers, rows in tables:
        content.extend(['<h2>' + html(name) + '</h2>', '<table cellspacing="0" cellpadding="5" border="1"><thead><tr>' + ''.join('<th>' + html(v) + '</th>' for v in headers) + '</tr></thead><tbody>'])
        for index, row in enumerate(rows): content.append('<tr bgcolor="' + ('#f5f1ed' if index % 2 == 0 else '#ffffff') + '">' + ''.join('<td>' + html(v) + '</td>' for v in row) + '</tr>')
        content.append('</tbody></table>')
    content.append('</body></html>'); return ''.join(content)


def write_pdf(path, title, info, tables):
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat); printer.setOutputFileName(str(path))
    printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4)); printer.setPageOrientation(QPageLayout.Orientation.Landscape)
    document = QTextDocument(); document.setDefaultFont(QFont('Arial', 9)); document.setHtml(report_html(title, info, tables)); document.print_(printer)
    if not Path(path).is_file() or Path(path).stat().st_size == 0: raise OSError('No se pudo crear el PDF.')


def breakdown_tables(rows):
    groups = OrderedDict()
    for row in rows: groups.setdefault((row.get('scene_id') or row['scene'], row['scene']), []).append(row)
    return [(heading, [ui_text('Categoría'), ui_text('Elemento'), ui_text('Estado'), ui_text('Texto de origen')], [[r['category'], r['element'], ui_text('Aprobado') if r['approved'] else ui_text('Pendiente'), r['source']] for r in elements]) for (identity, heading), elements in groups.items()] or [(ui_text('Desglose'), [ui_text('Estado')], [['No hay elementos.']])]


def budget_tables(rows, total, professional=True):
    keys = ['category', 'concept', 'quantity', 'unit', 'days', 'unit_price', 'currency', 'exchange', 'fringe', 'total', 'notes'] if professional else ['category', 'concept', 'quantity', 'days', 'unit_price', 'currency', 'total', 'notes']
    names = {'category': ui_text('Categoría'), 'concept': ui_text('Concepto'), 'quantity': ui_text('Cantidad'), 'unit': ui_text('Unidad'), 'days': ui_text('Jornadas'), 'unit_price': ui_text('Tarifa'), 'currency': ui_text('Moneda'), 'exchange': ui_text('Cambio'), 'fringe': 'Fringe %', 'total': 'Total base', 'notes': ui_text('Notas')}
    output = []
    for row in rows:
        values = [row.get(k, '') for k in keys]
        if row.get('level') != ui_text('Cuenta') and any(not row.get(k, '').strip() for k in ('quantity', 'days', 'unit_price', 'currency', 'exchange')):
            values[keys.index('total')] = ui_text('Pendiente')
        output.append(values)
    return [(ui_text('Presupuesto'), [names[k] for k in keys], output), ('Resumen', [ui_text('Concepto'), 'Importe'], [['Total calculado (revisar partidas pendientes)', total]])]


def schedule_tables(rows, selected=0):
    complete = (ui_text('Plan de rodaje'), [ui_text('Orden'), ui_text('Jornada'), ui_text('Fecha'), ui_text('Escena'), 'I/E', 'Luz', ui_text('Localización'), ui_text('Páginas'), ui_text('Sinopsis'), ui_text('Elementos')], [[r[k] for k in ('order', 'day', 'date', 'scene', 'interior', 'period', 'location', 'pages', 'synopsis', 'elements')] for r in rows])
    days = defaultdict(list); locations = defaultdict(list); elements = defaultdict(list)
    for row in rows:
        days[row['day'] or 'Sin jornada'].append(row); locations[row['location'] or 'Sin locación'].append(row)
        for name in row['elements'].split(','):
            if name.strip(): elements[name.strip()].append(row)
    daily = (ui_text('Necesidades por jornada'), [ui_text('Jornada'), ui_text('Escenas'), ui_text('Localizaciones'), ui_text('Elementos')], [[day, '\n'.join(r['scene'] for r in group), '\n'.join(sorted({r['location'] for r in group})), '\n'.join(sorted({r['elements'] for r in group}))] for day, group in days.items()])
    places = (ui_text('Escenas por localización'), [ui_text('Localización'), ui_text('Escenas'), ui_text('Jornadas')], [[place, '\n'.join(r['scene'] for r in group), ', '.join(sorted({r['day'] or 'Sin jornada' for r in group}))] for place, group in locations.items()])
    usage = (ui_text('Participación de elementos'), [ui_text('Elemento'), ui_text('Jornadas'), ui_text('Escenas')], [[element, ', '.join(sorted({r['day'] or 'Sin jornada' for r in group})), '\n'.join(r['scene'] for r in group)] for element, group in elements.items()])
    return [[complete], [daily], [places], [usage]][selected]
