"""Herramienta privada de QA compilada; excluida de la distribución de clientes."""
import json, os, sys
from pathlib import Path
from datetime import datetime, timezone, timedelta
from unittest.mock import patch
import suite_license as core
import suite_trial_state as state


def main():
    assert '__compiled__' in globals(), 'Esta validación requiere el ejecutable Nuitka.'
    out=Path(sys.argv[1]).resolve();out.mkdir(exist_ok=True)
    os.environ['QT_QPA_PLATFORM']='offscreen'
    os.environ['LOCALAPPDATA']=str(out/'local')
    perpetual=(out/'perpetual.ecplic').read_bytes();trial=(out/'trial.ecplic').read_bytes()
    first=datetime.now(timezone.utc);clock=[first];mirrors={}
    def read(key):return [mirrors[key]] if key in mirrors else []
    def write(key,value):mirrors[key]=value
    with patch.object(core,'license_path',return_value=out/'installed.json'), \
         patch.object(state,'state_directory',return_value=out/'activation'), \
         patch.object(state,'read_copies',side_effect=read),patch.object(state,'write_copies',side_effect=write), \
         patch.object(state,'now_utc',side_effect=lambda:clock[0]):
        core.install_license(perpetual);assert core.installed_license()['perpetual']
        core.install_license(trial);assert core.installed_license()['license_type']=='trial'
        clock[0]+=timedelta(days=2);assert core.installed_license()['duration_days']==3
        clock[0]=first+timedelta(days=3)
        try:core.installed_license()
        except state.TrialExpired:pass
        else:raise AssertionError('Compiled trial not expired')
        for key,value in [('license_type','perpetual'),('duration_days',30)]:
            changed=json.loads(trial);changed['payload'][key]=value
            try:core.verify(json.dumps(changed).encode())
            except ValueError:pass
            else:raise AssertionError('Compiled signature accepted tamper')
        core.install_license(perpetual)
    from PySide6.QtWidgets import QApplication,QFileDialog,QMessageBox
    from suite import SuiteWindow
    from project_store import fingerprint
    from openpyxl import load_workbook
    from pypdf import PdfReader
    app=QApplication([])
    QMessageBox.warning=lambda *a:(_ for _ in ()).throw(AssertionError(str(a)))
    QMessageBox.critical=QMessageBox.warning;QMessageBox.information=lambda *a:None
    w=SuiteWindow('professional');w.poll.stop();w.autosave.stop();w.show();app.processEvents()
    c=w.budget_controller;c.mutate('add','CARGO','500','QA')
    fringe=c.mutate_fringe('add',name='SEGURO',percentage='10')
    charge=c.mutate_charge('add',name='CONTRATO',amount='CARGO')
    credit=c.mutate_credit('add',name='DEVOLUCION',amount='1500')
    group=c.groups.mutate('add',name='UNIDAD')
    w.add_budget_row(data={'concept':'Compiled budget','quantity':'1','days':'1','unit_price':'10000',
                           'fringe_ids':[fringe],'charge_ids':[charge],'credit_ids':[credit],'group_ids':[group]})
    assert c.result.total==10000 and c.groups.analysis.totals[group]==10000
    w.modules.setCurrentWidget(w.budget_page);app.processEvents();w.grab().save(str(out/'compiled-budget.png'))
    w.path=out/'compiled-project.eguion';assert w.save();saved=w.project_data()
    for fmt in ['xlsx','pdf']:
        QFileDialog.getSaveFileName=lambda *a,f=fmt,**k:(str(out/('compiled-budget.'+f)),'')
        w.export_report('budget',fmt)
    book=load_workbook(out/'compiled-budget.xlsx');assert 'Contractual Charges' in book.sheetnames
    assert len(PdfReader(out/'compiled-budget.pdf').pages)>0
    w._last_saved=fingerprint(saved);w.dirty=False;w.close()
    w=SuiteWindow('professional');w.poll.stop();w.autosave.stop()
    assert w.open_path(out/'compiled-project.eguion')
    from time import perf_counter
    until=perf_counter()+30
    while getattr(w,'_import_job',None) is not None and perf_counter()<until:app.processEvents()
    assert getattr(w,'_import_job',None) is None and w.budget_controller.result.total==10000
    w._last_saved=fingerprint(w.project_data());w.dirty=False;w.close()
    (out/'compiled-results.json').write_text(json.dumps({'compiled_nuitka':True,'perpetual':True,'trial':True,
         'trial_reopen':True,'exact_expiry':True,'tamper_rejected':True,'budget_all_layers':True,
         'project_save_reopen':True,'excel':True,'pdf':True},indent=2))
    print('PASS: compiled licensing, UI, project save/reopen, all budget layers, PDF and Excel.')


if __name__=='__main__':main()
