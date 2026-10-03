from ui_i18n import ui_text, ui_join
from collections import OrderedDict
from datetime import datetime
from math import ceil
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.pagebreak import Break


def export_breakdown_workbook(filename, rows, title=ui_text('Sin título'), metadata=None):
    groups = OrderedDict()
    for row in rows:
        groups.setdefault((row.get('scene_id') or row.get('scene') or 'Sin escena', row.get('scene') or 'Sin escena'), []).append(row)
    book = Workbook()
    summary = book.active
    summary.title = 'Resumen de escenas'
    detail = book.create_sheet('Fichas de desglose')

    def write(sheet, row, values, fill=None, white=False, bold=False):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row, col, value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.font = Font(name='Calibri', size=11, bold=bold, color='FFFFFF' if white else '292524')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if fill:
                cell.fill = PatternFill('solid', fgColor=fill)

    def banner(sheet, row, text, columns, fill='7B2638'):
        write(sheet, row, [text] + [''] * (columns - 1), fill, True, True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
        sheet.row_dimensions[row].height = max(32, 18 * ceil(len(text) / 90))

    for sheet, widths in ((summary, [65, 14, 14, 65]), (detail, [32, 18, 95])):
        sheet.sheet_view.showGridLines = False
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + col)].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape' if sheet == summary else 'portrait'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.oddFooter.center.text = 'Eleuthera · Página &P de &N'
        banner(sheet, 1, title, len(widths))
        write(sheet, 2, ['Desglose · ' + datetime.now().strftime('%d/%m/%Y %H:%M')])
        sheet.freeze_panes = 'A4'
    write(summary, 3, [ui_text('Escena'), ui_text('Elementos'), 'Aprobados', 'Categorías'], '7B2638', True, True)
    summary.print_title_rows = '1:3'
    current = 4
    for index, ((scene_id, scene), elements) in enumerate(groups.items(), 4):
        categories = OrderedDict()
        for element in elements:
            categories.setdefault(element.get('category') or 'Sin categoría', []).append(element)
        category_text = ' · '.join(categories)
        write(summary, index, [scene, len(elements), sum(bool(item.get('approved')) for item in elements), category_text], 'F5F1ED' if index % 2 == 0 else 'FFFFFF')
        summary.row_dimensions[index].height = max(42, 16 * max(ceil(len(scene) / 55), ceil(len(category_text) / 55)))
        summary.cell(index, 1).hyperlink = f"#'Fichas de desglose'!A{current}"
        if current > 4:
            detail.row_breaks.append(Break(id=current - 1))
        banner(detail, current, scene, 3)
        current += 1
        write(detail, current, [ui_text('Elemento'), ui_text('Estado'), 'Descripción / texto de origen'], 'E9E2DB', bold=True)
        current += 1
        for category, items in categories.items():
            banner(detail, current, category, 3, '57534E')
            current += 1
            for item in items:
                values = [item.get('element', ''), ui_text('Aprobado') if item.get('approved') else ui_text('Pendiente'), item.get('source', '')]
                # Split very long evidence across rows to avoid Excel's row-height limit.
                chunks = []
                for paragraph in values[2].split('\n'):
                    chunks.extend(paragraph[i:i+75] for i in range(0, len(paragraph), 75))
                    if not paragraph:
                        chunks.append('')
                chunks = chunks or ['']
                for start in range(0, len(chunks), 18):
                    text = '\n'.join(chunks[start:start+18])
                    write(detail, current, [values[0] if start == 0 else '(continuación)', values[1] if start == 0 else '', text], 'FAF8F5' if current % 2 else 'FFFFFF')
                    detail.row_dimensions[current].height = max(32, 16 * max(len(chunks[start:start+18]), ceil(len(values[0])/28)) + 10)
                    current += 1
        current += 1
    if not groups:
        write(summary, 4, ['No hay elementos en el desglose.'])
        write(detail, 4, ['No hay elementos en el desglose.'])
    summary.auto_filter.ref = f'A3:D{max(3, summary.max_row)}'
    for sheet in (summary, detail):
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
    if metadata:
        from production_reports import metadata_lines, text_cell
        info = book.create_sheet(ui_text('Datos del proyecto'))
        text_cell(info.cell(1, 1), title)
        for index, line in enumerate(metadata_lines(metadata), 2): text_cell(info.cell(index, 1), line)
        info.column_dimensions['A'].width = 95
    book.save(filename)
