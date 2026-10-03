"""Plantillas profesionales de desarrollo estructural para Eleuthera Professional Plus.

Los nombres de metodologías sirven como referencias de organización narrativa. La
interfaz, campos y exportaciones de este módulo son diseños propios de Eleuthera.
"""
from __future__ import annotations
from ui_i18n import ui_text, ui_join

from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Iterable

ACCENT = "17616A"

COMMON_FIELDS = (
    ("purpose", "Propósito dramático"),
    ("character", "Personaje / foco"),
    ("objective", ui_text('Objetivo')),
    ("conflict", "Conflicto / tensión"),
    ("value", "Valor dramático"),
    ("value_start", "Valor inicial"),
    ("value_end", "Valor final"),
    ("notes", ui_text('Notas')),
)

# Campos adicionales sólo cuando aportan a la lógica de esa metodología.
EXTRA_FIELDS = {
    ui_text('Tres actos'): (("act", "Acto"), ("turn", "Giro / consecuencia")),
    ui_text('Cinco actos'): (("act", "Acto / movimiento"), ("escalation", "Escalada")),
    "Syd Field": (("act", "Acto"), ("plot_function", "Función del Plot Point")),
    "Save the Cat": (("transformation", "Transformación"),),
    "Freytag": (("movement", "Movimiento"), ("stakes", "Apuesta dramática")),
    "Viaje del héroe (Vogler)": (("world", "Mundo / umbral"), ("function", "Función arquetípica"), ("transformation", "Transformación")),
    "Story Circle (Dan Harmon)": (("step", "Paso"), ("cost", "Precio / costo"), ("transformation", ui_text('Cambio'))),
    ui_text('Ocho secuencias'): (("sequence_goal", "Objetivo de secuencia"), ("turn", "Giro de secuencia")),
    "McKee": (("gap", "Brecha expectativa/resultado"), ("stakes", "Apuesta"), ("turn", "Giro de valor")),
    "Kishōtenketsu": (("part_function", "Función de la parte"), ("contrast", "Contraste / relación")),
    "Truby — 22 pasos": (("revelation", "Revelación / decisión"), ("moral", "Dimensión moral"), ("transformation", "Transformación")),
    "Story Spine": (("cause", "Causa / consecuencia"), ("new_state", "Nuevo estado")),
    "Fichtean Curve": (("crisis", "Crisis"), ("stakes", "Aumento de apuesta")),
    "Seven-Point Story Structure": (("pressure", "Presión / giro"), ("setup_payoff", "Setup / payoff")),
    "Heroine’s Journey (Murdock)": (("identity", "Identidad / relación"), ("integration", "Integración"), ("transformation", "Transformación")),
    "Mini-Movie Method": (("segment_goal", "Objetivo del segmento"), ("hook", "Entrada / salida")),
    "TV — Teaser + 4 actos": (("act", "Bloque / acto"), ("act_out", "Act-out / gancho"), ("thread", "Trama A/B/C")),
    "TV — 6 actos": (("act", "Bloque / acto"), ("act_out", "Act-out / gancho"), ("thread", "Trama A/B/C")),
    "Círculo de conflicto": (("decision", "Decisión"), ("consequence", "Consecuencia")),
    "Arco de transformación": (("belief", "Creencia / carencia"), ("choice", "Elección"), ("transformation", "Transformación")),
    ui_text('Personalizado'): (("custom_function", "Función específica"),),
}

# Ubicaciones orientativas: describen la lógica propia de cada metodología.
# Sólo usamos porcentajes/páginas cuando el modelo realmente trabaja con ellos;
# en los demás casos mostramos fase, acto, secuencia o número de etapa.
SAVE_CAT_TARGETS = {
    "Imagen inicial": "≈ 1%", "Tema declarado": "≈ 5%", "Planteamiento": "≈ 1–10%",
    "Catalizador": "≈ 10%", "Debate": "≈ 10–20%", "Entrada al Acto II": "≈ 20%",
    "Trama B": "≈ 22%", "Diversión y juegos": "≈ 20–50%", "Midpoint": "≈ 50%",
    "Los malos se acercan": "≈ 50–75%", "Todo está perdido": "≈ 75%",
    "Noche oscura del alma": "≈ 75–80%", "Entrada al Acto III": "≈ 80%",
    "Finale": "≈ 80–99%", "Imagen final": "≈ 100%",
}

SYD_FIELD_TARGETS = {
    "Planteamiento": "Acto I · ≈ 0–25%", "Incidente incitador": "Acto I · temprano",
    "Plot Point I": "Fin Acto I · ≈ 17–25%", "Confrontación": "Acto II · ≈ 25–75%",
    "Midpoint": "Acto II · ≈ 50%", "Plot Point II": "Fin Acto II · ≈ 67–75%",
    "Resolución": "Acto III · ≈ 75–100%",
}

EIGHT_SEQUENCE_TARGETS = {
    f"Secuencia {i}: {name}": f"Secuencia {i}/8 · ≈ {(i-1)*12.5:g}–{i*12.5:g}%"
    for i, name in enumerate(("planteamiento", "detonación y giro", "primera progresión",
        "avance al punto medio", "consecuencias", "crisis y segundo giro",
        "preparación del clímax", "clímax y resolución"), 1)
}

def _ordered(beats, label="Etapa"):
    return {beat: f"{label} {i}/{len(beats)}" for i, beat in enumerate(beats, 1)}

ORIENTATION_MAPS = {
    ui_text('Tres actos'): {
        "Presentación":"Acto I · planteamiento", "Incidente detonante":"Acto I · ruptura",
        "Primer punto de giro":"Transición Acto I → II", "Midpoint":"Acto II · centro",
        "Crisis":"Acto II · tramo final", "Segundo punto de giro":"Transición Acto II → III",
        "Clímax":"Acto III · culminación", "Resolución":"Acto III · cierre"},
    ui_text('Cinco actos'): _ordered(("Exposición","Acción ascendente","Punto medio","Nueva complicación","Clímax","Acción descendente","Desenlace"), "Movimiento"),
    "Syd Field": SYD_FIELD_TARGETS,
    "Save the Cat": SAVE_CAT_TARGETS,
    "Freytag": _ordered(("Exposición","Acción ascendente","Clímax","Acción descendente","Resolución"), "Fase"),
    "Viaje del héroe (Vogler)": _ordered(("Mundo ordinario","Llamada a la aventura","Rechazo de la llamada","Encuentro con el mentor","Cruce del primer umbral","Pruebas, aliados y enemigos","Aproximación","Ordalía","Recompensa","Camino de regreso","Resurrección","Retorno con el elixir"), "Etapa"),
    "Story Circle (Dan Harmon)": _ordered(("Zona de confort",ui_text('Necesidad'),"Entrada a lo desconocido","Adaptación","Obtención","Precio","Regreso",ui_text('Cambio')), "Paso"),
    ui_text('Ocho secuencias'): EIGHT_SEQUENCE_TARGETS,
    "McKee": _ordered(("Incidente incitador","Complicaciones progresivas","Punto de no retorno","Crisis","Clímax","Resolución"), "Función"),
    "Kishōtenketsu": _ordered(("Ki — Introducción","Shō — Desarrollo","Ten — Giro","Ketsu — Conclusión"), "Parte"),
    "Truby — 22 pasos": _ordered(("Debilidad y necesidad","Fantasma / herida","Problema presente","Deseo","Aliado","Oponente","Falso aliado / oponente","Primera revelación y decisión","Plan","Plan del oponente","Impulso","Ataque del aliado","Derrota aparente","Segunda revelación y nueva decisión","Revelación para el público","Tercera revelación y decisión","Puerta / desafío final","Batalla","Autorrevelación","Decisión moral","Nuevo equilibrio","Estado final"), "Paso"),
    "Story Spine": _ordered(("Érase una vez","Cada día","Hasta que un día","A causa de eso I","A causa de eso II","Hasta que finalmente","Y desde entonces"), "Paso"),
    "Fichtean Curve": _ordered(("Crisis inicial","Crisis creciente I","Crisis creciente II","Crisis creciente III","Clímax","Acción descendente"), "Fase"),
    "Seven-Point Story Structure": _ordered(("Gancho","Primer punto de giro","Primer punto de presión","Punto medio","Segundo punto de presión","Segundo punto de giro","Resolución"), "Punto"),
    "Heroine’s Journey (Murdock)": _ordered(("Separación de lo femenino","Identificación con lo masculino","Camino de pruebas","Éxito aparente","Despertar espiritual","Descenso","Reconexión con lo femenino","Sanación de la ruptura","Integración"), "Etapa"),
    "Mini-Movie Method": _ordered(("Mini-película 1: planteamiento","Mini-película 2: nueva situación","Mini-película 3: progresión","Mini-película 4: punto medio","Mini-película 5: presión creciente","Mini-película 6: crisis","Mini-película 7: clímax","Mini-película 8: resolución"), "Segmento"),
    "TV — Teaser + 4 actos": _ordered(("Teaser / Cold open","Acto I","Primer giro / corte","Acto II","Midpoint / corte","Acto III","Crisis / corte","Acto IV","Clímax","Tag / cierre"), ui_text('Bloque')),
    "TV — 6 actos": _ordered(("Teaser / Cold open","Acto I","Acto II","Acto III","Midpoint","Acto IV","Acto V","Acto VI","Clímax","Tag / cierre"), ui_text('Bloque')),
    "Círculo de conflicto": _ordered(("Equilibrio inicial","Desequilibrio",ui_text('Objetivo'),"Obstáculo creciente","Decisión irreversible","Confrontación","Consecuencia","Nuevo equilibrio"), "Etapa"),
    "Arco de transformación": _ordered(("Estado inicial","Carencia interna","Desafío","Resistencia al cambio","Prueba decisiva","Crisis de identidad","Elección transformadora","Nuevo estado"), "Etapa"),
}

def orientation_for(model: str, beat: str) -> str:
    """Ubicación metodológica orientativa; nunca implica una regla obligatoria."""
    return ORIENTATION_MAPS.get(model, {}).get(beat, "")


def fields_for(model: str):
    """Devuelve columnas adaptadas. El núcleo común mantiene comparabilidad."""
    extras = EXTRA_FIELDS.get(model, ())
    # Los campos específicos van antes de las notas y el núcleo dramático.
    return tuple(extras) + COMMON_FIELDS


def ensure_development_store(dramatic_structure: dict) -> dict:
    store = dramatic_structure.setdefault("structure_development", {})
    return store if isinstance(store, dict) else {}


def row_data(dramatic_structure: dict, model: str, beat: str) -> dict:
    store = ensure_development_store(dramatic_structure)
    model_store = store.setdefault(model, {})
    row = model_store.setdefault(beat, {})
    return row


def value_change(row: dict) -> str:
    a, b = str(row.get("value_start", "")).strip(), str(row.get("value_end", "")).strip()
    return f"{a} → {b}" if a or b else ""


def export_xlsx(path: str, project_title: str, model: str, description: str,
                beats: Iterable[str], dramatic_structure: dict, scene_labels: dict[str, str]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Estructura"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    fields = fields_for(model)
    headers = ["Beat / etapa", "Ubicación orientativa", ui_text('Escena vinculada'), "Posición real"] + [label for _, label in fields] + ["Cambio de valor"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"] = project_title or "Proyecto sin título"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=ACCENT)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"] = f"Plantilla de desarrollo — {model}"
    ws["A2"].font = Font(size=12, bold=True, color=ACCENT)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    ws["A3"] = description
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 34
    thin = Side(style="thin", color="D9E1E3")
    for c, header in enumerate(headers, 1):
        cell = ws.cell(4, c, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
    points = dramatic_structure.get("structure_points", {}) or {}
    beat_list = list(beats)
    for r, beat in enumerate(beat_list, 5):
        data = row_data(dramatic_structure, model, beat)
        sid = str(points.get(beat, ""))
        scene = scene_labels.get(sid, "")
        pos = ""
        if scene and scene_labels:
            ids = list(scene_labels.keys())
            try: pos = f"{(ids.index(sid)+1)/len(ids)*100:.1f}%"
            except ValueError: pass
        vals = [beat, orientation_for(model, beat), scene, pos] + [data.get(key, "") for key, _ in fields] + [value_change(data)]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(r, c, val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
        ws.row_dimensions[r].height = 34
    widths = [28, 34, 14] + [24] * len(fields) + [18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, 4+len(beat_list))}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    summary = wb.create_sheet("Resumen")
    summary.sheet_view.showGridLines = False
    summary["A1"] = "ELEUTHERA — RESUMEN ESTRUCTURAL"
    summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=ACCENT)
    summary.merge_cells("A1:D1")
    rows = [(ui_text('Proyecto'), project_title or "Proyecto sin título"), ("Modelo", model),
            ("Hitos / etapas", len(beat_list)),
            ("Hitos con escena", sum(1 for b in beat_list if points.get(b))),
            ("Hitos desarrollados", sum(1 for b in beat_list if any(str(v).strip() for v in row_data(dramatic_structure, model, b).values())))]
    for i, (k, v) in enumerate(rows, 3):
        summary.cell(i, 1, k).font = Font(bold=True, color=ACCENT)
        summary.cell(i, 2, v)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 48
    summary["A10"] = "Nota"
    summary["A10"].font = Font(bold=True, color=ACCENT)
    summary["B10"] = "Las posiciones y metodologías son herramientas de desarrollo, no reglas obligatorias."
    summary["B10"].alignment = Alignment(wrap_text=True)
    wb.save(path)


def export_pdf(path: str, project_title: str, model: str, description: str,
               beats: Iterable[str], dramatic_structure: dict, scene_labels: dict[str, str]):
    """PDF profesional usando el motor Qt ya distribuido con Eleuthera."""
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    fields = fields_for(model)
    points = dramatic_structure.get("structure_points", {}) or {}
    body = [
        "<html><head><style>",
        "body{font-family:Arial,sans-serif;color:#202426;font-size:9pt}",
        "h1{color:#17616A;font-size:19pt;margin-bottom:2px} h2{font-size:12pt;color:#17616A}",
        "table{border-collapse:collapse;width:100%;} th{background:#17616A;color:white;padding:5px;font-size:8pt}",
        "td{border:1px solid #d9e1e3;padding:5px;vertical-align:top;font-size:8pt} .muted{color:#5f6b6d}",
        "</style></head><body>",
        f"<h1>{escape(project_title or 'Proyecto sin título')}</h1>",
        f"<h2>Plantilla de desarrollo — {escape(model)}</h2>",
        f"<p class='muted'>{escape(description)}</p>",
        "<table><tr><th>Beat / etapa</th><th>Escena</th><th>Desarrollo</th><th>Cambio de valor</th></tr>",
    ]
    for beat in beats:
        data = row_data(dramatic_structure, model, beat)
        sid = str(points.get(beat, ""))
        chunks = []
        for key, label in fields:
            val = str(data.get(key, "")).strip()
            if val:
                chunks.append(f"<b>{escape(label)}:</b> {escape(val)}")
        body.append("<tr>" +
                    f"<td><b>{escape(beat)}</b></td>" +
                    f"<td>{escape(scene_labels.get(sid, ''))}</td>" +
                    f"<td>{'<br/>'.join(chunks) or '—'}</td>" +
                    f"<td>{escape(value_change(data) or '—')}</td></tr>")
    body.append("</table><p class='muted'>Generado con Eleuthera Cinema Suite Professional Plus.</p></body></html>")
    doc = QTextDocument()
    doc.setHtml("".join(body))
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(str(path))
    printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    printer.setPageOrientation(QPageLayout.Orientation.Landscape)
    printer.setPageMargins(QMarginsF(10, 10, 10, 10), QPageLayout.Unit.Millimeter)
    doc.print_(printer)
