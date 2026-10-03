from ui_i18n import ui_text, ui_join

import re

from collections import defaultdict

from datetime import date



from PySide6.QtCore import QDate, Qt

from PySide6.QtGui import QBrush, QColor, QFont

from eleuthera_spreadsheet import StructuredSpreadsheet

from PySide6.QtWidgets import (

    QAbstractItemView, QComboBox, QDateEdit, QFileDialog, QFrame, QHBoxLayout,

    QApplication, QColorDialog, QHeaderView, QInputDialog, QLabel, QMenu, QMessageBox, QPushButton, QSizePolicy, QSpinBox, QTableWidget,

    QTableWidgetItem, QTabWidget, QVBoxLayout,

)





STRIP_COLORS = {

    ('INT', 'DÍA'): '#f4e8a5',

    ('INT', 'NOCHE'): '#b9c8e8',

    ('EXT', 'DÍA'): '#d7efd0',

    ('EXT', 'NOCHE'): '#c7b8dc',

}





def item(value='', editable=True):

    cell = QTableWidgetItem(str(value))

    if not editable:

        cell.setFlags(cell.flags() & ~Qt.ItemFlag.ItemIsEditable)

    return cell





class StripboardTable(StructuredSpreadsheet):

    def __init__(self, parent=None):

        super().__init__(0, 11, parent)

        self.setHorizontalHeaderLabels((ui_text('Orden'), ui_text('Jornada'), ui_text('Fecha'), ui_text('Escena'), 'I/E', 'D/N', ui_text('Localización'), ui_text('Páginas'), ui_text('Sinopsis'), ui_text('Elementos'), ui_text('Estado')))

        self.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)

        self.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)

        self.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)

        self.setDragDropOverwriteMode(False)

        self.setAlternatingRowColors(False)


        # Rendimiento: no usar ResizeToContents en una stripboard grande.

        # Ese modo vuelve a medir miles de celdas cada vez que cambia la tabla.

        header = self.horizontalHeader()

        for column in range(self.columnCount()):

            header.setSectionResizeMode(column, QHeaderView.ResizeMode.Interactive)

        header.setSectionResizeMode(8, QHeaderView.ResizeMode.Stretch)

        header.setSectionResizeMode(9, QHeaderView.ResizeMode.Stretch)

        for column, width in {0: 64, 1: 82, 2: 96, 3: 270, 4: 54, 5: 82, 6: 220, 7: 76, 10: 110}.items():

            header.resizeSection(column, width)

        self.setWordWrap(False)

        self.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Fixed)

        self.verticalHeader().setDefaultSectionSize(26)

        self._drag_resize_columns = tuple(column for column in range(self.columnCount()) if column not in (8, 9))



    def _selected_rows(self):

        rows = sorted({item.row() for item in self.selectedItems()})
        if not rows and self.currentRow() >= 0:
            rows = [self.currentRow()]
        return rows



    def startDrag(self, supported_actions):

        # Las columnas ya trabajan con anchos fijos/interactivos, por lo que Qt

        # no necesita recalcular su contenido durante el arrastre.

        super().startDrag(supported_actions)



    def dropEvent(self, event):

        rows = self._selected_rows()

        if not rows:

            return super().dropEvent(event)

        target = self.rowAt(event.position().toPoint().y())

        if target < 0:

            target = self.rowCount()



        self.setUpdatesEnabled(False)

        try:

            records = [[self.item(row, column).clone() for column in range(self.columnCount())] for row in rows]

            for row in reversed(rows):

                self.removeRow(row)

                if row < target:

                    target -= 1

            for offset, record in enumerate(records):

                self.insertRow(target + offset)

                for column, cell in enumerate(record):

                    self.setItem(target + offset, column, cell)

            self.clearSelection()

            self.selectRow(target)

        finally:

            self.setUpdatesEnabled(True)

            self.viewport().update()



        event.accept()

        if hasattr(self.parent(), 'renumber'):

            self.parent().renumber(refresh_dates=False)





class SchedulingPage(QFrame):

    def __init__(self, parent=None):

        super().__init__(parent); self.setObjectName('workspace')

        layout = QVBoxLayout(self); layout.setContentsMargins(30, 25, 30, 28)

        tools = QHBoxLayout()

        self.sync_button = QPushButton(ui_text('Actualizar…')); self.sync_button.setObjectName('primary')

        self.sync_button.setToolTip(ui_text('Revisar actualizaciones desde el guion y el desglose.'))

        self.day = QSpinBox(); self.day.setRange(1, 999); self.day.setPrefix(ui_text('Jornada: '))

        assign = QPushButton(ui_text('Asignar a seleccionadas')); assign.clicked.connect(self.assign_day)

        up = QPushButton(ui_text('Subir')); up.clicked.connect(lambda: self.move_selected(-1))

        down = QPushButton(ui_text('Bajar')); down.clicked.connect(lambda: self.move_selected(1))

        self.start_date = QDateEdit(QDate.currentDate()); self.start_date.setCalendarPopup(True); self.start_date.setDisplayFormat('dd/MM/yyyy')

        self.start_date.dateChanged.connect(self.refresh_dates)

        self.report_type = QComboBox(); self.report_type.addItems((ui_text('Plan completo'), ui_text('Necesidades por jornada'), ui_text('Escenas por localización'), ui_text('Participación de elementos')))

        export = QPushButton('Excel'); export.clicked.connect(self.export_reports)

        for widget in (self.sync_button, QLabel(ui_text('Inicio:')), self.start_date, self.day, assign, up, down): tools.addWidget(widget)

        tools.addStretch(); layout.addLayout(tools)

        self.report_actions = QHBoxLayout()

        self.report_actions.addWidget(self.report_type); self.report_actions.addWidget(export); self.report_actions.addStretch()

        layout.addLayout(self.report_actions)

        self.table = StripboardTable(self)
        self.table.itemChanged.connect(self.on_item_changed)

        # PLAN DE RODAJE · herramientas tipo Excel agrupadas y compactas.
        excel_tools = QHBoxLayout()
        self.plan_excel_tools = excel_tools
        edit_btn = QPushButton(ui_text('Edición') + ' ▾'); edit_menu = QMenu(edit_btn)
        edit_menu.addAction(ui_text('Cortar'), self.plan_cut); edit_menu.addAction(ui_text('Copiar'), self.plan_copy); edit_menu.addAction(ui_text('Pegar'), self.plan_paste); edit_btn.setMenu(edit_menu)
        format_label = QLabel(ui_text('Formato:'))
        bold_btn = QPushButton('B'); bold_btn.setCheckable(True); bold_btn.setFixedWidth(32); bold_btn.clicked.connect(lambda checked: self.plan_format_font('bold', checked))
        italic_btn = QPushButton('I'); italic_btn.setCheckable(True); italic_btn.setFixedWidth(32); italic_btn.clicked.connect(lambda checked: self.plan_format_font('italic', checked))
        underline_btn = QPushButton('U'); underline_btn.setCheckable(True); underline_btn.setFixedWidth(32); underline_btn.clicked.connect(lambda checked: self.plan_format_font('underline', checked))
        self.plan_font_size = QComboBox(); self.plan_font_size.addItems(tuple(str(n) for n in (8, 9, 10, 11, 12, 14, 16, 18, 20, 24))); self.plan_font_size.setCurrentText('10'); self.plan_font_size.setFixedWidth(72); self.plan_font_size.currentTextChanged.connect(self.plan_format_font_size)
        text_btn = QPushButton(ui_text('Texto')); text_btn.clicked.connect(self.plan_format_text_color)
        bg_btn = QPushButton(ui_text('Fondo')); bg_btn.clicked.connect(self.plan_format_background)
        border_btn = QPushButton(ui_text('Bordes')); border_btn.clicked.connect(self.plan_format_borders)
        align_btn = QPushButton(ui_text('Alinear') + ' ▾'); align_menu = QMenu(align_btn)
        align_menu.addAction(ui_text('Izquierda'), lambda: self.plan_format_alignment(Qt.AlignmentFlag.AlignLeft)); align_menu.addAction(ui_text('Centro'), lambda: self.plan_format_alignment(Qt.AlignmentFlag.AlignHCenter)); align_menu.addAction(ui_text('Derecha'), lambda: self.plan_format_alignment(Qt.AlignmentFlag.AlignRight)); align_btn.setMenu(align_menu)
        clear_btn = QPushButton(ui_text('Limpiar')); clear_btn.clicked.connect(self.plan_clear_format)
        search_btn = QPushButton(ui_text('Buscar') + ' ▾'); search_menu = QMenu(search_btn)
        search_menu.addAction(ui_text('Buscar'), self.plan_find); search_menu.addAction(ui_text('Buscar y reemplazar'), self.plan_replace); search_btn.setMenu(search_menu)
        data_btn = QPushButton(ui_text('Datos') + ' ▾'); data_menu = QMenu(data_btn)
        data_menu.addAction(ui_text('Orden ascendente'), lambda: self.plan_sort(True)); data_menu.addAction(ui_text('Orden descendente'), lambda: self.plan_sort(False)); data_menu.addSeparator(); data_menu.addAction(ui_text('Filtrar'), self.plan_filter); data_menu.addAction(ui_text('Mostrar todo'), self.plan_show_all); data_btn.setMenu(data_menu)
        for widget in (edit_btn, format_label, bold_btn, italic_btn, underline_btn, self.plan_font_size, text_btn, bg_btn, border_btn, align_btn, clear_btn, search_btn, data_btn): excel_tools.addWidget(widget)
        excel_tools.addStretch(); layout.addLayout(excel_tools)
        layout.addWidget(self.table, 1)

        self.table.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)

        self._plan_maximized = False

        self._plan_max_tabbar = None

        self._plan_max_tabbar_was_visible = True

        self._plan_max_margins = layout.contentsMargins()

        self._plan_max_vscroll_policy = self.table.verticalScrollBarPolicy()

        self._plan_max_hscroll_policy = self.table.horizontalScrollBarPolicy()

        footer = QHBoxLayout(); self.summary = QLabel(ui_text('0 escenas · 0 jornadas')); self.summary.setObjectName('moduleHint')

        self.legend = QLabel(ui_text('INT/DÍA amarillo · INT/NOCHE azul · EXT/DÍA verde · EXT/NOCHE violeta')); self.legend.setObjectName('moduleHint')

        footer.addWidget(self.summary); footer.addStretch(); footer.addWidget(self.legend); layout.addLayout(footer)





    def _plan_selected_cells(self):
        return self.table.selected_items_including_empty() if hasattr(self, 'table') else []

    def plan_format_background(self):
        cells = self._plan_selected_cells()
        if not cells: return
        initial = cells[0].background().color() if cells[0].background().style() != Qt.BrushStyle.NoBrush else QColor('#ffffff')
        color = QColorDialog.getColor(initial, self, ui_text('Fondo de celda'))
        if color.isValid():
            for cell in cells: cell.setBackground(QBrush(color))

    def plan_format_text_color(self):
        cells = self._plan_selected_cells()
        if not cells: return
        initial = cells[0].foreground().color() if cells[0].foreground().style() != Qt.BrushStyle.NoBrush else QColor('#111111')
        color = QColorDialog.getColor(initial, self, ui_text('Color del texto'))
        if color.isValid():
            for cell in cells: cell.setForeground(QBrush(color))

    def plan_format_font(self, kind, enabled):
        for cell in self._plan_selected_cells():
            font = cell.font()
            if kind == 'bold': font.setBold(bool(enabled))
            elif kind == 'italic': font.setItalic(bool(enabled))
            elif kind == 'underline': font.setUnderline(bool(enabled))
            cell.setFont(font)

    def plan_format_font_size(self, value):
        try: size = int(value)
        except (TypeError, ValueError): return
        for cell in self._plan_selected_cells():
            font = cell.font(); font.setPointSize(size); cell.setFont(font)

    def plan_format_alignment(self, horizontal):
        for cell in self._plan_selected_cells():
            cell.setTextAlignment(horizontal | Qt.AlignmentFlag.AlignVCenter)

    def plan_format_borders(self):
        from eleuthera_spreadsheet import BORDER_ROLE
        cells = self._plan_selected_cells()
        if not cells: return
        color = QColorDialog.getColor(QColor('#707070'), self, ui_text('Color de borde'))
        if not color.isValid(): return
        for cell in cells: cell.setData(BORDER_ROLE, color.name())
        self.table.viewport().update()

    def plan_clear_format(self):
        from eleuthera_spreadsheet import BORDER_ROLE
        rows = set()
        for cell in self._plan_selected_cells():
            rows.add(cell.row())
            cell.setFont(QFont())
            cell.setTextAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
            cell.setData(BORDER_ROLE, None)
        # En stripboard el color base codifica INT/EXT y DÍA/NOCHE: restaurarlo,
        # no blanquearlo, al limpiar formato manual.
        blocked = self.table.blockSignals(True)
        try:
            for row in rows: self.color_strip(row)
        finally:
            self.table.blockSignals(blocked)
        self.table.viewport().update()

    def plan_copy(self):
        ranges = self.table.selectedRanges()
        if not ranges: return
        rg = ranges[0]; lines = []
        for row in range(rg.topRow(), rg.bottomRow() + 1):
            lines.append('\t'.join(self.table.item(row, col).text() if self.table.item(row, col) else ''
                                   for col in range(rg.leftColumn(), rg.rightColumn() + 1)))
        QApplication.clipboard().setText('\n'.join(lines))

    def plan_cut(self):
        self.plan_copy()
        for cell in self._plan_selected_cells():
            if cell.flags() & Qt.ItemFlag.ItemIsEditable: cell.setText('')

    def plan_paste(self):
        text = QApplication.clipboard().text()
        if not text: return
        sr, sc = self.table.currentRow(), self.table.currentColumn()
        if sr < 0 or sc < 0: return
        matrix = [row.split('\t') for row in text.rstrip('\n').splitlines()]
        for rr, row in enumerate(matrix):
            for cc, value in enumerate(row):
                r, c = sr + rr, sc + cc
                if r >= self.table.rowCount() or c >= self.table.columnCount(): continue
                cell = self.table.ensure_item(r, c)
                if cell.flags() & Qt.ItemFlag.ItemIsEditable: cell.setText(value)

    def plan_find(self):
        query, ok = QInputDialog.getText(self, ui_text('Buscar'), ui_text('Texto a buscar:'))
        if not ok or not query: return
        columns = self.table.columnCount(); total = self.table.rowCount() * columns
        if not total: return
        start = max(0, self.table.currentRow() * columns + self.table.currentColumn() + 1)
        for offset in range(total):
            pos = (start + offset) % total; row, col = divmod(pos, columns)
            cell = self.table.item(row, col)
            if cell and query.casefold() in cell.text().casefold():
                self.table.setCurrentCell(row, col); self.table.scrollToItem(cell); return
        QMessageBox.information(self, ui_text('Buscar'), ui_text('No se encontraron coincidencias.'))

    def plan_replace(self):
        find, ok = QInputDialog.getText(self, ui_text('Buscar y reemplazar'), ui_text('Buscar:'))
        if not ok or not find: return
        replacement, ok = QInputDialog.getText(self, ui_text('Buscar y reemplazar'), ui_text('Reemplazar por:'))
        if not ok: return
        changed = 0
        for cell in self._plan_selected_cells():
            if cell.flags() & Qt.ItemFlag.ItemIsEditable and find.casefold() in cell.text().casefold():
                cell.setText(re.sub(re.escape(find), lambda m: replacement, cell.text(), flags=re.IGNORECASE)); changed += 1
        QMessageBox.information(self, ui_text('Buscar y reemplazar'), ui_text('Reemplazos: ') + str(changed))

    def plan_sort(self, ascending=True):
        column = self.table.currentColumn()
        if column < 0: return
        self.table.sortItems(column, Qt.SortOrder.AscendingOrder if ascending else Qt.SortOrder.DescendingOrder)
        self.renumber(refresh_dates=False)

    def plan_filter(self):
        query, ok = QInputDialog.getText(self, ui_text('Filtrar plan'), ui_text('Mostrar filas que contengan:'))
        if not ok: return
        needle = query.casefold().strip()
        for row in range(self.table.rowCount()):
            haystack = ' '.join(self.table.item(row, col).text() if self.table.item(row, col) else ''
                               for col in range(self.table.columnCount())).casefold()
            self.table.setRowHidden(row, bool(needle) and needle not in haystack)

    def plan_show_all(self):
        for row in range(self.table.rowCount()): self.table.setRowHidden(row, False)


    def _suite_tab_bar(self):

        """Devuelve la barra principal que contiene Plan de rodaje, si existe."""

        widget = self.parentWidget()

        while widget is not None:

            if isinstance(widget, QTabWidget) and widget.indexOf(self) >= 0:

                return widget.tabBar()

            widget = widget.parentWidget()

        return None



    def toggle_plan_maximized(self):

        """Maximiza/restaura la tabla sin alterar datos ni lógica del plan."""

        self.set_plan_maximized(not self._plan_maximized)



    def set_plan_maximized(self, maximized):

        maximized = bool(maximized)

        if maximized == self._plan_maximized:

            return

        self._plan_maximized = maximized



        button = getattr(self, 'maximize_plan_button', None)

        layout = self.layout()

        if maximized:

            self._plan_max_margins = layout.contentsMargins()

            self._plan_max_vscroll_policy = self.table.verticalScrollBarPolicy()

            self._plan_max_hscroll_policy = self.table.horizontalScrollBarPolicy()

            self._plan_max_tabbar = self._suite_tab_bar()

            if self._plan_max_tabbar is not None:

                self._plan_max_tabbar_was_visible = self._plan_max_tabbar.isVisible()

                self._plan_max_tabbar.hide()



            layout.setContentsMargins(4, 4, 4, 4)

            self.table.setMinimumHeight(0)

            self.table.setMaximumHeight(16777215)

            self.table.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOn)

            self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOn)

            if button is not None:

                button.setText(ui_text('↙ Restaurar'))

                button.setToolTip('Restaurar la vista normal del Plan de rodaje.')

        else:

            layout.setContentsMargins(self._plan_max_margins)

            self.table.setVerticalScrollBarPolicy(self._plan_max_vscroll_policy)

            self.table.setHorizontalScrollBarPolicy(self._plan_max_hscroll_policy)

            if self._plan_max_tabbar is not None and self._plan_max_tabbar_was_visible:

                self._plan_max_tabbar.show()

            self._plan_max_tabbar = None

            if button is not None:

                button.setText(ui_text('⛶ Maximizar plan'))

                button.setToolTip('Maximizar la tabla del Plan de rodaje.')



        self.table.updateGeometry()

        self.updateGeometry()



    @staticmethod

    def scene_info(heading):

        heading = re.sub(r"^\d+[A-Z]?[.\-:)]?\s+(?=(?:INT|EXT|I/E))", "", heading.strip(), flags=re.I)

        text = heading.upper().strip()

        interior = 'EXT' if text.startswith(('EXT.', 'EXT ', 'I/E')) else 'INT'

        period = 'NOCHE' if any(word in text for word in ('NOCHE', 'ANOCHECER', 'MADRUGADA')) else 'DÍA'

        location = re.sub(r'^(INT\.?|EXT\.?|INT\.?/EXT\.?|I/E)\s*', '', heading, flags=re.I)

        location = re.split(r'\s+-\s+(?:D[IÍ]A|NOCHE|TARDE|AMANECER|ANOCHECER|MADRUGADA)\b', location, flags=re.I)[0].strip()

        return interior, period, location



    def sync_from_project(self, blocks, breakdown):

        old = {self.table.item(row, 3).text(): self.row_data(row) for row in range(self.table.rowCount())}

        scenes = []; current = None

        for block in blocks:

            if block.get('type') == 'scene':

                current = {'heading': block.get('text', '').strip(), 'body': []}; scenes.append(current)

            elif current and block.get('text', '').strip(): current['body'].append(block.get('text', '').strip())

        elements = defaultdict(list)

        for row in breakdown:

            value = row.get('element', '').strip()

            if value and value not in elements[row.get('scene', '')]: elements[row.get('scene', '')].append(value)



        # Carga masiva real: preasigna todas las filas y no ejecuta geometría

        # de la tabla hasta terminar. Esto evita insertRow() + relayout N veces.

        previous_signals = self.table.blockSignals(True)

        previous_updates = self.table.updatesEnabled()

        self.table.setUpdatesEnabled(False)

        try:

            self.table.setRowCount(len(scenes))

            for row, scene in enumerate(scenes):

                number = row + 1

                heading = scene['heading']; previous = old.get(heading, {})

                interior, period, location = self.scene_info(heading)

                synopsis = previous.get('synopsis') or (' '.join(scene['body'])[:180])

                pages = previous.get('pages') or max(0.125, round((len(scene['body']) + 1) / 8, 3))

                self._set_strip_row(row, {

                    'order': number, 'day': previous.get('day', 1), 'scene': heading,

                    'interior': previous.get('interior', interior), 'period': previous.get('period', period),

                    'location': previous.get('location', location), 'pages': pages, 'synopsis': synopsis,

                    'elements': ', '.join(elements.get(heading, [])), 'status': previous.get('status', ui_text('Pendiente')),

                })

            self.renumber()

        finally:

            self.table.blockSignals(previous_signals)

            self.table.setUpdatesEnabled(previous_updates)

            if previous_updates:

                self.table.viewport().update()



    def _set_strip_row(self, row, data):

        values = (data.get('order', row + 1), data.get('day', ''), '', data.get('scene', ''), data.get('interior', 'INT'), data.get('period', 'DÍA'), data.get('location', ''), data.get('pages', 0.125), data.get('synopsis', ''), data.get('elements', ''), data.get('status', ui_text('Pendiente')))

        for column, value in enumerate(values):

            self.table.setItem(row, column, item(value, editable=column not in (0, 2)))

        self.table.item(row, 0).setData(Qt.ItemDataRole.UserRole, data.get('scene_id', ''))

        self.color_strip(row)



    def add_strip(self, data):

        row = self.table.rowCount()

        self.table.insertRow(row)

        self._set_strip_row(row, data)



    def row_data(self, row):

        keys = ('order', 'day', 'date', 'scene', 'interior', 'period', 'location', 'pages', 'synopsis', 'elements', 'status')

        return {**{key: self.table.item(row, column).text() for column, key in enumerate(keys)}, 'scene_id': self.table.item(row, 0).data(Qt.ItemDataRole.UserRole) or ''}



    def to_data(self): return [self.row_data(row) for row in range(self.table.rowCount())]



    def load_data(self, rows, start_date=None):

        if start_date: self.start_date.setDate(QDate.fromString(start_date, Qt.DateFormat.ISODate))



        rows = list(rows or [])

        previous_signals = self.table.blockSignals(True)

        previous_updates = self.table.updatesEnabled()

        self.table.setUpdatesEnabled(False)

        try:

            self.table.setRowCount(len(rows))

            for row_index, row_data in enumerate(rows):

                self._set_strip_row(row_index, row_data)

            self.renumber()

        finally:

            self.table.blockSignals(previous_signals)

            self.table.setUpdatesEnabled(previous_updates)

            if previous_updates:

                self.table.viewport().update()



    def renumber(self, refresh_dates=True):

        self.table.blockSignals(True)

        for row in range(self.table.rowCount()):

            self.table.item(row, 0).setText(str(row + 1))

        self.table.blockSignals(False)

        if refresh_dates:

            self.refresh_dates()



    def refresh_dates(self):

        # Solo Jornada/Fecha. No recolorea toda la tabla: cambiar una fecha no

        # modifica INT/EXT ni DÍA/NOCHE, y ese repintado era el cuello de botella.

        self.table.blockSignals(True)

        days = set()

        try:

            for row in range(self.table.rowCount()):

                day_cell = self.table.item(row, 1)

                date_cell = self.table.item(row, 2)

                if date_cell is None:

                    date_cell = item('', editable=False)

                    self.table.setItem(row, 2, date_cell)



                raw_text = day_cell.text().strip() if day_cell is not None else ''

                # Compatibilidad con planes heredados: una fila sin jornada se
                # considera inicialmente Jornada 1, igual que las escenas nuevas.
                # Así cambiar 'Inicio' actualiza inmediatamente la columna Fecha.
                if not raw_text and day_cell is not None:
                    day_cell.setText('1')
                    raw_text = '1'

                try:

                    raw = float(raw_text)

                    if not raw.is_integer() or raw < 1:

                        raise ValueError()

                    shooting_day = int(raw)

                except (ValueError, TypeError):

                    if date_cell.text():

                        date_cell.setText('')

                    continue



                days.add(shooting_day)

                value = self.start_date.date().addDays(shooting_day - 1).toString('dd/MM/yyyy')

                if date_cell.text() != value:

                    date_cell.setText(value)

        finally:

            self.table.blockSignals(False)

        self.summary.setText(ui_text('{} escenas · {} jornadas').format(self.table.rowCount(), len(days)))



    def color_strip(self, row):

        interior_cell = self.table.item(row, 4)

        period_cell = self.table.item(row, 5)

        interior = interior_cell.text().upper() if interior_cell is not None else 'INT'

        period = period_cell.text().upper() if period_cell is not None else 'DÍA'

        color = QColor(STRIP_COLORS.get((interior, period), '#eeeeee'))

        text = QBrush(QColor('#252329'))

        for column in range(self.table.columnCount()):

            cell = self.table.item(row, column)

            if cell is not None:

                cell.setBackground(color)

                cell.setForeground(text)



    def on_item_changed(self, cell):

        if cell.column() == 1:

            self.refresh_dates()

        elif cell.column() in (4, 5):

            # I/E y D/N solo cambian el color de ESTA fila.

            self.table.blockSignals(True)

            try:

                self.color_strip(cell.row())

            finally:

                self.table.blockSignals(False)



    def assign_day(self):

        rows = self.table._selected_rows()

        for row in rows: self.table.item(row, 1).setText(str(self.day.value()))

        self.refresh_dates()



    def move_selected(self, direction):

        rows = self.table._selected_rows()

        if len(rows) != 1: return

        row = rows[0]; target = row + direction

        if target < 0 or target >= self.table.rowCount(): return

        self.table.blockSignals(True)

        first = [self.table.takeItem(row, col) for col in range(self.table.columnCount())]

        second = [self.table.takeItem(target, col) for col in range(self.table.columnCount())]

        for col in range(self.table.columnCount()): self.table.setItem(row, col, second[col]); self.table.setItem(target, col, first[col])

        self.table.blockSignals(False); self.table.selectRow(target); self.renumber(refresh_dates=False)



    def export_reports(self):

        if hasattr(self, 'export_callback'):

            return self.export_callback()

        filename, _ = QFileDialog.getSaveFileName(self, ui_text('Exportar informes de producción'), '', 'Excel (*.xlsx)')

        if not filename: return

        if not filename.lower().endswith('.xlsx'): filename += '.xlsx'

        try:

            from openpyxl import Workbook

            book = Workbook(); plan = book.active; plan.title = ui_text('Plan de rodaje')

            plan.append([self.table.horizontalHeaderItem(col).text() for col in range(self.table.columnCount())])

            for row in self.to_data(): plan.append(list(row.values()))

            by_day = book.create_sheet(ui_text('Necesidades por jornada')); by_day.append((ui_text('Jornada'), ui_text('Fecha'), ui_text('Escenas'), ui_text('Localizaciones'), ui_text('Elementos')))

            grouped = defaultdict(lambda: {'dates': set(), 'scenes': [], 'locations': set(), 'elements': set()})

            for row in self.to_data():

                group = grouped[row['day']]; group['dates'].add(row['date']); group['scenes'].append(row['scene']); group['locations'].add(row['location'])

                group['elements'].update(value.strip() for value in row['elements'].split(',') if value.strip())

            for day in sorted(grouped, key=lambda value: int(float(value))):

                group = grouped[day]; by_day.append((day, ', '.join(group['dates']), '\n'.join(group['scenes']), ', '.join(sorted(group['locations'])), ', '.join(sorted(group['elements']))))

            elements = book.create_sheet('Participación elementos'); elements.append((ui_text('Elemento'), ui_text('Jornadas'), ui_text('Escenas')))

            usage = defaultdict(lambda: {'days': set(), 'scenes': []})

            for row in self.to_data():

                for value in (part.strip() for part in row['elements'].split(',') if part.strip()): usage[value]['days'].add(row['day']); usage[value]['scenes'].append(row['scene'])

            for name, values in sorted(usage.items()): elements.append((name, ', '.join(sorted(values['days'])), '\n'.join(values['scenes'])))

            book.save(filename); QMessageBox.information(self, ui_text('Informes de producción'), ui_text('Plan e informes exportados correctamente.'))

        except Exception as error: QMessageBox.critical(self, ui_text('Informes de producción'), str(error))

