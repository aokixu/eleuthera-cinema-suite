import copy,json,re
from pathlib import Path
from time import perf_counter
from dataclasses import asdict
from decimal import Decimal
from PySide6.QtCore import Qt,QTimer
from PySide6.QtWidgets import QApplication,QDialog,QLineEdit,QDialogButtonBox,QMessageBox,QFileDialog
from openpyxl import load_workbook
from pypdf import PdfReader
from suite import SuiteWindow
from project_store import fingerprint
from budget_values import BudgetError
from budget_credits_ui import CreditsDialog,CreditAssignmentDialog,CreditDetailDialog

app=QApplication([]);out=Path('test-results/stage5');out.mkdir(exist_ok=True)
passed,metrics,warnings=[],{},[]
QMessageBox.warning=lambda *a,**k:warnings.append(a)
QMessageBox.information=lambda *a,**k:None
QMessageBox.question=lambda *a,**k:QMessageBox.StandardButton.Yes
QMessageBox.critical=lambda *a,**k:(_ for _ in ()).throw(AssertionError(a))
def window():
    w=SuiteWindow('professional');w.poll.stop();w.autosave.stop();return w
def close(w):w._last_saved=fingerprint(w.project_data());w.dirty=False;w.close()
def drain(predicate):
    deadline=perf_counter()+90
    while predicate() and perf_counter()<deadline:app.processEvents()
    assert not predicate()
def update(c,identifier,**changes):
    fields=asdict(c.credits.get(identifier));fields.pop('id');fields.update(changes)
    return c.mutate_credit('edit',identifier,**fields)
def rejects(w,call):
    before=fingerprint(w.project_data())
    try:call()
    except BudgetError:pass
    else:raise AssertionError('Invalid operation accepted')
    assert fingerprint(w.project_data())==before

w=window();c=w.budget_controller
c.mutate('add','APORTE','1500','Proveedor');c.mutate('add','OTRO_APORTE','1500','')
dialog=CreditsDialog(c,w);old_exec=QDialog.exec;fields=('PROVEEDOR','APORTE','Devolución proveedor')
def fill(form):
    for field,text in zip(form.findChildren(QLineEdit),fields):field.setText(text)
    form.findChild(QDialogButtonBox).accepted.emit();return form.result()
QDialog.exec=fill;dialog.edit(True);credit=c.credits.items()[0].id
before=fingerprint(w.project_data());dialog.edit(True);assert warnings and fingerprint(w.project_data())==before
fields=('PROVEEDOR_RENOMBRADO','APORTE','Devuelve importe');dialog.edit(False)
assert c.credits.get(credit).name=='PROVEEDOR_RENOMBRADO'
QDialog.exec=old_exec
w.add_budget_row(data={'concept':'Costo original','quantity':'1','days':'1','unit_price':'10000'})
assignment=CreditAssignmentDialog(c,0,w);assignment.listing.item(0).setCheckState(Qt.CheckState.Checked);assignment.save()
assert w.budget.item(0,12).text()=='8500.00' and w.budget.item(0,8).text()=='10000'
bonus=c.mutate_credit('add',name='BONO',amount='250.5')
c.assign_credits([0],[credit,bonus,credit]);assert c.result.total==Decimal('8249.5')
assert w.budget_data()[0]['credit_ids']==[credit,bonus]
detail=CreditDetailDialog(c,0,w);detail.show();app.processEvents();detail.grab().save(str(out/'credit-details.png'))
assert detail.table.rowCount()==2;detail.close()
dialog.table.selectRow(0);dialog.toggle();assert c.result.total==Decimal('9749.5')
dialog.toggle();dialog.remove();assert 'fila 1' in str(warnings[-1])
c.assign_credits([0],[]);assert c.result.total==10000
c.assign_credits([0],[credit]);assert c.result.total==8500
unused=c.mutate_credit('add',name='TEMPORAL',amount='1')
dialog.refresh();dialog.table.selectRow(2);dialog.remove();assert len(c.credits.items())==2
dialog.show();app.processEvents();dialog.grab().save(str(out/'credits.png'));dialog.close()
passed.append('UI crear/editar/renombrar/activar/eliminar/asignar/desasignar; detalle múltiple y deduplicación sin alterar tarifa')

rejects(w,lambda:update(c,credit,amount='-1'))
rejects(w,lambda:update(c,credit,amount='1/0'))
rejects(w,lambda:update(c,credit,amount='NO_EXISTE'))
rejects(w,lambda:c.assign_credits([0,999],[credit]))
rejects(w,lambda:c.assign_credits([0],['missing']))
rejects(w,lambda:c.mutate('remove','APORTE'))
rejects(w,lambda:c.mutate('edit','APORTE','CAMBIADO','1500',''))
rejects(w,lambda:c.mutate('edit','APORTE','APORTE','-1',''))
w.budget.item(0,0).setText('Subdetalle');before=fingerprint(w.project_data())
w.budget.item(0,0).setText('Cuenta')
assert w.budget.item(0,0).text()=='Subdetalle' and fingerprint(w.project_data())==before
passed.append('Operaciones rechazadas atómicas; referencia de Global protegida; conversión a Cuenta restaurada sin Groups')

update(c,credit,amount='OTRO_APORTE')
c.mutate('edit','APORTE','APORTE','1400','');assert c.last_run['evaluated_rows']==0
c.mutate('edit','OTRO_APORTE','OTRO_APORTE','1600','');assert c.result.total==8400
update(c,credit,amount='APORTE')
c.mutate('edit','APORTE','APORTE','10','')
update(c,bonus,amount='5')
temp=c.mutate_credit('add',name='SOLO_TEMPORAL',amount='1')
w.add_budget_row(data={'concept':'Se elimina','unit_price':'10','credit_ids':[temp]})
w.budget.clearSelection();w.budget.selectRow(1);w.remove_budget_rows();c.mutate_credit('remove',temp)
passed.append('Cambio de dependencias con importe idéntico; eliminar una partida libera referencias de Credit')

fringe=c.mutate_fringe('add',name='SEGURO',percentage='10')
a=c.groups.mutate('add',name='UNIDAD_A');b=c.groups.mutate('add',name='UNIDAD_B')
c.mutate_credit('add',name='INACTIVO',amount='99',active=False)
w._loading=True;w.budget.setRowCount(0)
w.add_budget_row(data={'level':'Cuenta','concept':'Equipo','block':'ATL'})
for i in range(1025):
    w.add_budget_row(data={'concept':'Partida '+str(i),'block':'ATL','quantity':'1','days':'1','unit_price':'1000',
                          'fringe_ids':[fringe],'credit_ids':[credit] if i<1000 else [bonus],
                          'group_ids':[a] if i<25 else [b] if i>=1000 else []})
w._loading=False;c.request();drain(lambda:c.busy)
assert c.result.total==Decimal(1117375)
before=c.result.totals[:1001];start=perf_counter();update(c,bonus,amount='6');drain(lambda:c.busy)
metrics['credit_25_rows_ms']=(perf_counter()-start)*1000
assert c.last_run['evaluated_rows']==25 and c.result.totals[1:1001]==before[1:]
assert c.groups.analysis.changed=={b}
metrics['rare_evaluated_rows']=c.last_run['evaluated_rows']
start=perf_counter();w.budget.item(1,5).setText('2');drain(lambda:c.busy)
metrics['edit_one_row_ms']=(perf_counter()-start)*1000
assert c.last_run['evaluated_rows']==1 and c.groups.analysis.changed=={a}
assert w.budget.item(1,12).text()=='2190.00'
w.budget.item(1,5).setText('1')
passed.append('1025 partidas: Credit de 25 recalcula solo 25 y un Group; editar una partida recalcula una')

ticks=[perf_counter()];timer=QTimer();timer.setInterval(1);timer.timeout.connect(lambda:ticks.append(perf_counter()));timer.start()
start=perf_counter();update(c,credit,amount='APORTE + 1')
metrics['credit_1000_submit_ms']=(perf_counter()-start)*1000
assert c.busy;drain(lambda:c.busy);metrics['credit_1000_complete_ms']=(perf_counter()-start)*1000
timer.stop();assert len(ticks)>2 and c.last_run['evaluated_rows']==1000
metrics['gui_heartbeat_events']=len(ticks)-1;metrics['max_gui_gap_ms']=max(y-x for x,y in zip(ticks,ticks[1:]))*1000
start=perf_counter();c.mutate('edit','APORTE','APORTE','12','');drain(lambda:c.busy)
metrics['global_credit_1000_ms']=(perf_counter()-start)*1000
assert c.result.total==Decimal(1114350) and c.groups.analysis.totals[a]==27175
passed.append('Credit y Global sobre 1000 partidas: Qt responde, Fringes preceden a Credits y Groups reciben el neto')

c.groups.set_filter([a]);project=out/'credits-project.eguion';w.path=project
start=perf_counter();assert w.save();metrics['save_1026_rows_ms']=(perf_counter()-start)*1000
saved=w.project_data();assert len(saved['budget'])==1026;close(w)
w=window();c=w.budget_controller
start=perf_counter();assert w.open_path(project);drain(lambda:getattr(w,'_import_job',None) is not None)
metrics['reopen_1026_rows_ms']=(perf_counter()-start)*1000
assert w.project_data()['budget_credits']==saved['budget_credits']
assert c.result.total==Decimal(1114350) and not c.credits.items()[-1].active
assert w.budget_data()[1]['credit_ids']==[credit]
for mode in ('missing','account','negative','global'):
    bad=copy.deepcopy(saved)
    if mode=='missing':bad['budget'][1]['credit_ids']=['missing']
    if mode=='account':bad['budget'][0]['credit_ids']=[credit]
    if mode=='negative':bad['budget_credits']['items'][0]['amount']='-1'
    if mode=='global':bad['budget_credits']['items'][0]['amount']='NO_EXISTE'
    rejects(w,lambda:w.load_project(bad))
passed.append('Guardar filtrado y reabrir conserva todo, incluidos IDs/estado; cargas corruptas no reemplazan el proyecto')

saved_bytes=project.read_bytes();update(c,credit,amount='APORTE + 2');assert c.busy
assert not w.save() and project.read_bytes()==saved_bytes
for fmt in ('xlsx','pdf'):
    pending=out/('pending.'+fmt);QFileDialog.getSaveFileName=lambda *args,**kwargs:(str(pending),'')
    w.export_report('budget',fmt);assert not pending.exists()
drain(lambda:c.busy);assert c.result.total==Decimal(1113350)
excel=out/'credits-budget.xlsx';QFileDialog.getSaveFileName=lambda *args,**kwargs:(str(excel),'')
start=perf_counter();w.export_report('budget','xlsx');metrics['export_excel_ms']=(perf_counter()-start)*1000
book=load_workbook(excel);assert 'Credits' in book.sheetnames and 'Base y Fringes' in book.sheetnames
values=[str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row if value is not None]
assert '1086.00' in values and '14.00' in values and '1100.00' in values
assert w.total_label.text() in values
passed.append('Guardar/exportar pendientes protegidos; Excel conserva formatos previos y añade desglose Credits actualizado')

update(c,credit,amount='15');assert c.busy
w.budget.item(1,5).setText('2');drain(lambda:c.busy)
assert c.last_run['mode']=='full' and w.budget.item(1,12).text()=='2185.00'
passed.append('Edición durante recálculo cancela resultados obsoletos de Credits y Groups')

small=copy.deepcopy(w.project_data());small['budget']=small['budget'][1:2]
small['budget'][0].update(quantity='1',unit_price='10000',exchange='6.96',currency='USD')
small['budget_credits']['items'][0]['amount']='1500'
small.pop('budget_currencies',None);small.pop('budget_currency_legacy',None)
small['project_info'].update(currency='BOB',rates={'USD':6.96});small['budget_settings']['base_currency']='BOB'
w.load_project(small)
assert c.result.before_credits[0]==76560 and c.result.credit_totals[0]==10440 and c.result.total==66120
assert c.groups.analysis.totals[a]==66120
pdf=out/'credits-budget.pdf';QFileDialog.getSaveFileName=lambda *args,**kwargs:(str(pdf),'')
start=perf_counter();w.export_report('budget','pdf');metrics['export_small_pdf_ms']=(perf_counter()-start)*1000
text=''.join(page.extract_text() or '' for page in PdfReader(pdf).pages)
assert re.search(r'66120\s*\.\s*00',text) and re.search(r'10440\s*\.\s*00',text)
small.pop('budget_credits');small['budget'][0].pop('credit_ids');w.load_project(small)
assert c.result.total==76560 and not c.credits.items()
passed.append('Conversión real BOB sin duplicación, PDF con antes/reducción/neto y proyecto antiguo sin Credits')
close(w)
(out/'integration-results.json').write_text(json.dumps(dict(passed=passed,metrics=metrics),ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS:',len(passed),'escenarios de Credits');print(json.dumps(metrics,indent=2))
