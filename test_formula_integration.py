"""Etapa 2: Qt, persistencia, informes, invalidación y métricas reproducibles."""
import copy
import json
from pathlib import Path
import tempfile
from time import perf_counter
from decimal import Decimal
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QLineEdit, QMessageBox, QFileDialog
from openpyxl import load_workbook
from suite import SuiteWindow
from budget_model import Global, Globals, BudgetError
from budget_globals_ui import GlobalsDialog
from project_store import fingerprint

app = QApplication([])
output = Path('test-results/stage2')
output.mkdir(exist_ok=True)
metrics, passed, warnings = {}, [], []
QMessageBox.warning = lambda *a, **k: warnings.append(a)
QMessageBox.information = lambda *a, **k: None
QMessageBox.question = lambda *a, **k: QMessageBox.StandardButton.Yes
QMessageBox.critical = lambda *a, **k: (_ for _ in ()).throw(AssertionError(a))

def window():
    w = SuiteWindow('professional'); w.poll.stop(); w.autosave.stop()
    return w

def drain(predicate):
    end = perf_counter() + 60
    while predicate() and perf_counter() < end: app.processEvents()
    assert not predicate(), 'Operación no terminó'

def close(w):
    w._last_saved = fingerprint(w.project_data()); w.dirty = False; w.close()

w = window(); c = w.budget_controller
for name, value in [('DIAS_PREPRO', '5'), ('DIAS_TOTAL', 'DIAS_RODAJE + DIAS_PREPRO'), ('SEMANAS_RODAJE', 'DIAS_RODAJE / 5'), ('NUM_EXTRAS', '20'), ('TARIFA_BASE', '100'), ('TARIFA_EXTRA', 'TARIFA_BASE * 1.5')]:
    c.mutate('add', name, value, '')
c.mutate('edit', 'DIAS_RODAJE', 'DIAS_RODAJE', '15', 'Rodaje')
assert c.catalog.get('DIAS_TOTAL').value == 20
assert c.catalog.get('SEMANAS_RODAJE').value == 3
passed.append('Global constante, calculado y múltiples dependencias')

dialog = GlobalsDialog(c, w)
original_exec = QDialog.exec
fields = ('DIAS_RODAJE', 'DIAS_TOTAL + 1', 'No aceptar')
def fill(form):
    for widget, value in zip(form.findChildren(QLineEdit), fields): widget.setText(value)
    form.findChild(QDialogButtonBox).accepted.emit()
    return form.result()
QDialog.exec = fill
before = fingerprint(w.project_data())
dialog.table.selectRow(0); dialog.edit(False)
assert warnings and 'circular' in str(warnings[-1]).lower()
assert fingerprint(w.project_data()) == before
fields = ('TEMP_CALC', 'TARIFA_EXTRA / 2', 'Calculado desde el formulario')
dialog.edit(True)
assert c.catalog.get('TEMP_CALC').value == 75
dialog.refresh()
assert dialog.table.columnCount() == 4
assert dialog.table.item(dialog.table.rowCount()-1, 1).text().startswith('ƒ')
assert dialog.table.item(dialog.table.rowCount()-1, 2).text() == '75.0'
dialog.show(); app.processEvents()
dialog.grab().save(str(output / 'globals-formulas.png'))
dialog.table.selectRow(dialog.table.rowCount()-1); dialog.remove()
assert 'TEMP_CALC' not in c.catalog.values()
QDialog.exec = original_exec
dialog.close()
passed.append('UI fórmula/resultado y rechazo atómico de ciclo')

# Bidirectional compatibility does not overwrite a calculated DIAS_RODAJE.
c.mutate('edit', 'DIAS_RODAJE', 'DIAS_RODAJE', 'DIAS_PREPRO + 10', '')
assert not w.global_days.isEnabled()
w.global_days.setValue(1)
assert c.catalog.get('DIAS_RODAJE').value == 15
c.mutate('edit', 'DIAS_RODAJE', 'DIAS_RODAJE', '15', '')
c.mutate('add', 'INVERSO', '1 / DIAS_RODAJE', '')
before = fingerprint(w.project_data())
w.global_days.setValue(0)
assert c.catalog.get('DIAS_RODAJE').value == 15
assert w.global_days.value() == 15
assert fingerprint(w.project_data()) == before
c.mutate('remove', 'INVERSO')
passed.append('Compatibilidad del control Días rodaje y error sin excepción Qt')

w._loading = True
w.add_budget_row(data={'level': 'Cuenta', 'concept': 'Figuración', 'block': 'ATL'})
for i in range(1001):
    w.add_budget_row(data={'concept': 'Extra ' + str(i), 'block': 'ATL', 'quantity': 'NUM_EXTRAS', 'days': 'DIAS_TOTAL', 'unit_price': 'TARIFA_EXTRA'})
w.add_budget_row(data={'level': 'Cuenta', 'concept': 'Independientes', 'block': 'BTL'})
for i in range(1000):
    w.add_budget_row(data={'concept': 'Fijo ' + str(i), 'quantity': '2', 'days': '1', 'unit_price': '50'})
w._loading = False
c.request(); drain(lambda: c.busy)
assert c.result.total == Decimal(60160000)
independent = c.result.totals[1002:]
ticks = [perf_counter()]
timer = QTimer(); timer.setInterval(1); timer.timeout.connect(lambda: ticks.append(perf_counter())); timer.start()
start = perf_counter()
c.mutate('edit', 'DIAS_RODAJE', 'DIAS_RODAJE', '16', '')
metrics['update_1001_submit_ms'] = (perf_counter()-start)*1000
assert c.busy
drain(lambda: c.busy)
metrics['update_1001_complete_ms'] = (perf_counter()-start)*1000
timer.stop()
metrics['gui_heartbeat_events'] = len(ticks)-1
metrics['max_gui_gap_ms'] = max(b-a for a, b in zip(ticks, ticks[1:]))*1000
assert len(ticks) > 2
assert c.result.total == Decimal(63163000)
assert c.result.totals[1002:] == independent
assert c.last_run == {'mode': 'incremental', 'evaluated_rows': 1001, 'updated_rows': 1002}
metrics['selective_recalculation'] = dict(c.last_run)
passed.append('Propagación a 1001 partidas: 1000 independientes sin recalcular; cuentas y totales')

c.mutate('edit', 'DIAS_TOTAL', 'DIAS_TOTAL', 'DIAS_RODAJE + DIAS_PREPRO', 'Descripción actualizada')
assert c.last_run['evaluated_rows'] == 0
assert c.last_run['mode'] == 'unchanged'
try: c.mutate('remove', 'DIAS_RODAJE')
except BudgetError: pass
else: raise AssertionError('Eliminó un Global referenciado')
try: c.mutate('edit', 'TARIFA_EXTRA', 'OTRA_TARIFA', 'TARIFA_BASE * 1.5', '')
except BudgetError: pass
else: raise AssertionError('Renombró un Global de partidas')
passed.append('Edición de descripción sin recálculo y protección de referencias')

with tempfile.TemporaryDirectory() as folder:
    folder = Path(folder)
    project = folder / 'formulas.eguion'; report = folder / 'presupuesto.xlsx'
    w.path = project
    start = perf_counter(); assert w.save(); metrics['save_2003_rows_ms'] = (perf_counter()-start)*1000
    saved_bytes = project.read_bytes()
    data = json.loads(saved_bytes)
    expected = copy.deepcopy(data['budget_globals'])
    assert any(x.get('formula') == 'DIAS_RODAJE + DIAS_PREPRO' for x in expected['items'])
    assert all('value' not in x for x in expected['items'] if x['kind'] == 'formula')
    # Loading invalid dependencies or cycles does not replace any live state.
    for formula in ('DIAS_TOTAL', 'NO_EXISTE'):
        invalid = copy.deepcopy(data)
        invalid['budget_globals']['items'][0] = {'name': 'DIAS_RODAJE', 'kind': 'formula', 'formula': formula}
        before = fingerprint(w.project_data())
        try: w.load_project(invalid)
        except BudgetError: pass
        else: raise AssertionError('Carga inválida aceptada')
        assert fingerprint(w.project_data()) == before
    close(w)
    w = window(); c = w.budget_controller
    start = perf_counter(); assert w.open_path(project)
    drain(lambda: getattr(w, '_import_job', None) is not None)
    metrics['open_2003_rows_async_ms'] = (perf_counter()-start)*1000
    assert w.project_data()['budget_globals'] == expected
    assert c.result.total == Decimal(63163000)
    assert w.budget.item(1, 5).text() == 'NUM_EXTRAS'
    assert w.budget.item(1, 7).text() == 'DIAS_TOTAL'
    assert w.budget.item(1, 8).text() == 'TARIFA_EXTRA'
    passed.append('Guardar, cerrar y reabrir realmente: reconstruye grafo y resultados')
    # No save or export of partially updated totals.
    c.mutate('edit', 'DIAS_RODAJE', 'DIAS_RODAJE', '17', '')
    assert c.busy
    assert not w.save() and project.read_bytes() == saved_bytes
    QFileDialog.getSaveFileName = lambda *a, **k: (str(report), '')
    w.export_report('budget', 'xlsx'); assert not report.exists()
    drain(lambda: c.busy)
    assert c.result.total == Decimal(66166000)
    assert not w.export_issues('budget')
    w.export_report('budget', 'xlsx')
    book = load_workbook(report)
    values = [str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row if value is not None]
    assert w.total_label.text() in values
    assert '66000.00' in values
    assert w.save()
    passed.append('Guardar/exportar durante pendiente y Excel con resultados nuevos')
    # Cancelled incremental work must never publish over a later row edit.
    c.mutate('edit', 'DIAS_RODAJE', 'DIAS_RODAJE', '18', '')
    w.budget.item(1, 5).setText('NUM_EXTRAS * 2')
    drain(lambda: c.busy)
    assert c.last_run['mode'] == 'full'
    assert w.budget.item(1, 12).text() == '138000.00'
    passed.append('Edición durante recálculo cancela el resultado obsoleto')
    # Stage 1 and pre-Globals schemas and the legacy token remain supported.
    w.load_project({'blocks': [], 'budget_settings': {'shooting_days': 12},
                    'budget': [{'quantity': '2', 'days': '@DIAS_RODAJE', 'unit_price': '100', 'currency': 'USD', 'exchange': '1', 'fringe': '0'}]})
    assert w.budget.item(0, 12).text() == '2400.00'
    w.global_days.setValue(15)
    assert w.budget.item(0, 12).text() == '3000.00'
    old = w.project_data(); old['budget_globals'] = {'version': 1, 'items': [{'name': 'DIAS_RODAJE', 'value': '10'}]}
    w.load_project(old)
    assert w.budget.item(0, 12).text() == '2000.00'
    passed.append('Proyectos numéricos antiguos, Globals v1 y @DIAS_RODAJE')
close(w)

chain = [Global('G0', 1)] + [Global('G'+str(i), 'G'+str(i-1)+' + 1') for i in range(1, 100)]
start = perf_counter(); model = Globals(chain); metrics['chain_100_load_ms'] = (perf_counter()-start)*1000
start = perf_counter(); model.edit('G0', 'G0', '2'); metrics['chain_100_propagate_ms'] = (perf_counter()-start)*1000
assert model.get('G99').value == 101
assert len(model.last_recalculated) == 99
passed.append('Cadena de 100 Globals sin recursión del grafo')
(output/'integration-results.json').write_text(json.dumps({'passed': passed, 'metrics': metrics}, ensure_ascii=False, indent=2), encoding='utf-8')
print('PASS:', len(passed), 'escenarios de integración de Etapa 2')
print(json.dumps(metrics, indent=2))
